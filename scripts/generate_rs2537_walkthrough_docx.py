from __future__ import annotations

from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUT_DIR = Path("generated")
OUT_DIR.mkdir(parents=True, exist_ok=True)
OUT_PATH = OUT_DIR / "ITU-R_RS.2537-0_detailed_calculation_walkthrough_CN.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, east_asia: str = "宋体", latin: str = "Aptos", size: float | None = None,
                 bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_spacing(p, before=0, after=0, line=1.5) -> None:
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    fmt.line_spacing = line


def add_text(doc: Document, text: str, *, bold: bool = False, align=None,
             size: float = 10.5, color: str | None = None, indent: bool = True,
             before: float = 0, after: float = 4) -> None:
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    set_paragraph_spacing(p, before=before, after=after)
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, color=color)


def add_bullet(doc: Document, text: str, level: int = 0) -> None:
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.left_indent = Cm(0.74 + 0.55 * level)
    p.paragraph_format.first_line_indent = Cm(-0.35)
    set_paragraph_spacing(p, after=2)
    r = p.add_run(text)
    set_run_font(r, size=10.5)


def add_numbered(doc: Document, text: str, level: int = 0) -> None:
    style = "List Number" if level == 0 else "List Number 2"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.left_indent = Cm(0.74 + 0.55 * level)
    p.paragraph_format.first_line_indent = Cm(-0.35)
    set_paragraph_spacing(p, after=2)
    r = p.add_run(text)
    set_run_font(r, size=10.5)


def add_equation(doc: Document, lines: str | Iterable[str], note: str | None = None) -> None:
    if isinstance(lines, str):
        lines = [lines]
    box = doc.add_table(rows=1, cols=1)
    box.alignment = WD_TABLE_ALIGNMENT.CENTER
    box.autofit = True
    cell = box.cell(0, 0)
    set_cell_shading(cell, "F4F7FB")
    set_cell_margins(cell, top=100, start=180, bottom=100, end=180)
    for idx, line in enumerate(lines):
        p = cell.paragraphs[0] if idx == 0 else cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(p, after=0, line=1.0)
        r = p.add_run(line)
        set_run_font(r, east_asia="Cambria Math", latin="Cambria Math", size=11.5)
    if note:
        p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(p, before=2, after=0, line=1.0)
        r = p.add_run(note)
        set_run_font(r, size=9, color="566573")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_heading(text, level=level)
    set_paragraph_spacing(p, before=10 if level == 1 else 6, after=5, line=1.0)
    for r in p.runs:
        set_run_font(
            r,
            east_asia="微软雅黑" if level <= 2 else "黑体",
            latin="Aptos Display",
            size={1: 16, 2: 13, 3: 11.5}.get(level, 10.5),
            bold=True,
            color="17365D" if level <= 2 else "1F4E79",
        )


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False if widths else True
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        set_cell_shading(hdr[i], "D9EAF7")
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(hdr[i])
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(p, after=0, line=1.0)
        r = p.add_run(h)
        set_run_font(r, east_asia="微软雅黑", size=9.5, bold=True, color="17365D")
        if widths:
            hdr[i].width = Cm(widths[i])
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[i])
            if ridx % 2 == 1:
                set_cell_shading(cells[i], "F8FAFC")
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i > 0 else WD_ALIGN_PARAGRAPH.LEFT
            set_paragraph_spacing(p, after=0, line=1.0)
            r = p.add_run(str(value))
            set_run_font(r, size=9.3)
            if widths:
                cells[i].width = Cm(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color="666666")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    run2 = paragraph.add_run(" 页")
    set_run_font(run2, size=9, color="666666")


def add_toc(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("目录")
    set_run_font(r, east_asia="微软雅黑", size=16, bold=True, color="17365D")
    p = doc.add_paragraph()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = 'TOC \\o "1-3" \\h \\z \\u'
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "separate")
    fld_char3 = OxmlElement("w:t")
    fld_char3.text = "在Word中右击此处并选择“更新域”，即可生成目录。"
    fld_char4 = OxmlElement("w:fldChar")
    fld_char4.set(qn("w:fldCharType"), "end")
    r = p.add_run()
    r._r.append(fld_char)
    r._r.append(instr_text)
    r._r.append(fld_char2)
    r._r.append(fld_char3)
    r._r.append(fld_char4)
    set_run_font(r, size=10.5, color="666666")


def create_document() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.1)
    sec.bottom_margin = Cm(1.9)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.0)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)

    for style_name, east, size, color in (
        ("Title", "微软雅黑", 22, "17365D"),
        ("Subtitle", "宋体", 11, "566573"),
        ("Heading 1", "微软雅黑", 16, "17365D"),
        ("Heading 2", "微软雅黑", 13, "17365D"),
        ("Heading 3", "黑体", 11.5, "1F4E79"),
    ):
        st = styles[style_name]
        st.font.name = "Aptos Display"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), east)
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(color)

    header = sec.header.paragraphs[0]
    header.text = "ITU-R RS.2537-0 方法详细演算"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for r in header.runs:
        set_run_font(r, east_asia="微软雅黑", size=9, bold=True, color="5B7083")

    footer = sec.footer.paragraphs[0]
    add_page_number(footer)

    # Cover page
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(54)
    r = p.add_run("ITU-R RS.2537-0 方法详细演算")
    set_run_font(r, east_asia="微软雅黑", latin="Aptos Display", size=24, bold=True, color="17365D")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    r = p.add_run("SAR5 对 GPS SBAS 地面参考接收机的脉冲干扰初步评估")
    set_run_font(r, east_asia="微软雅黑", latin="Aptos Display", size=16, bold=True, color="1F4E79")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(20)
    r = p.add_run("基于 Report ITU-R RS.2537-0（2023）附件2典型算例")
    set_run_font(r, size=11, color="566573")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(72)
    r = p.add_run("选取工况：1236.5 MHz / 14 MHz Chirp / PRF 1860 Hz / 脉宽 37.63 μs")
    set_run_font(r, east_asia="微软雅黑", size=11, bold=True, color="2F5597")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(110)
    r = p.add_run("中文工程解读与逐步复算")
    set_run_font(r, size=11, color="666666")

    doc.add_page_break()
    add_toc(doc)
    doc.add_page_break()

    add_heading(doc, "1  文档说明与案例选择", 1)
    add_text(doc, "RS.2537-0在ITU-R分类中属于研究报告（Report），并非建议书（Recommendation）。该报告给出了如何将ITU-R M.2030的脉冲干扰评估模型应用于“1215–1300 MHz星载合成孔径雷达（SAR）干扰卫星无线电导航业务（RNSS）地面接收机”的初步评估实例。")
    add_text(doc, "本文件完整复算报告附件2中的一组代表性参数：标准SAR5以1236.5 MHz为中心频率、14 MHz线性调频带宽、PRF为1860 Hz、脉宽为37.63 μs，照射GPS卫星增强系统（SBAS）地面参考接收机。该工况是报告表13中14 MHz波形对应的最大PDC工况，报告给出的最终新增劣化为0.387 dB。")
    add_text(doc, "本文中的公式、参数和判定逻辑均以RS.2537-0为基础；文末关于允许PDC和允许重叠带宽的反算，是在报告公式基础上的工程延伸，不是报告直接列出的数值。", bold=True, color="9C5700")

    add_heading(doc, "2  RS.2537-0方法的总体流程", 1)
    add_table(
        doc,
        ["阶段", "核心问题", "主要输出"],
        [
            ["1. 峰值链路计算", "最坏主瓣照射时，RNSS输入端峰值功率多大？", "P_RFI,pk"],
            ["2. 生存与饱和判定", "是否烧毁？是否进入压缩或饱和？", "生存裕量、饱和超量"],
            ["3. 频谱重叠计算", "SAR Chirp中有多少频率落入RNSS通带？", "Δf、重叠比η"],
            ["4. 有效脉宽与PDC", "每个脉冲使RNSS实际丢失多长时间？", "PW_eff、PDC_Y"],
            ["5. 多源聚合", "新增SAR与既有脉冲背景如何叠加？", "PDC_base+Y、R_I,total"],
            ["6. 性能劣化判定", "新增源造成的有效噪声劣化是否超限？", "D_dB与允许值比较"],
        ],
        widths=[3.6, 8.0, 4.5],
    )
    add_text(doc, "报告将评估分成两层：第一层进行前端安全与饱和筛查，假定RNSS接收机处在SAR主波束内，暂不计接收机滤波；第二层在确认“不会损坏但可能饱和”后，进一步计入调频谱与接收通带的重叠、接收机恢复时间、PDC以及基线干扰环境。")
    add_text(doc, "若峰值功率低于接收机生存电平，但高于“输入饱和电平以下15 dB”的工程筛查线，则应继续计算接收机性能劣化。", bold=True, color="C00000")

    add_heading(doc, "3  本次演算使用的参数", 1)
    add_heading(doc, "3.1  SAR5发射参数", 2)
    add_table(
        doc,
        ["参数", "取值", "说明"],
        [
            ["峰值发射功率", "6120 W（37.9 dBW）", "报告表1"],
            ["发射天线峰值增益", "36.6 dBi", "主瓣最坏几何"],
            ["峰值EIRP", "74.5 dBW", "37.9 + 36.6"],
            ["中心频率", "1236.5 MHz", "SAR5可选中心频率之一"],
            ["线性调频带宽", "14 MHz", "Chirp带宽"],
            ["脉冲宽度", "37.63 μs", "与PRF联动以保持约7%占空比"],
            ["PRF", "1860 Hz", "表13高PRF工况"],
            ["发射占空比", "约7%", "37.63 μs × 1860 Hz"],
            ["最坏星地距离", "636.65 km", "报告表11"],
            ["自由空间损耗", "150.5 dB", "报告表11"],
            ["极化失配损耗", "1.46 dB", "报告采用值"],
        ],
        widths=[5.0, 4.3, 7.0],
    )

    add_heading(doc, "3.2  SBAS地面参考接收机参数", 2)
    add_table(
        doc,
        ["参数", "取值", "物理意义"],
        [
            ["接收天线增益", "−2 dBi", "最坏峰值链路预算采用值"],
            ["输入饱和电平", "−135 dBW", "超过后发生增益压缩/饱和"],
            ["输入生存电平", "−10 dBW", "峰值超过可能损伤前端"],
            ["过载恢复时间 τ_r", "1 μs", "脉冲结束后恢复所需时间"],
            ["预相关滤波器中心", "1227.6 MHz", "GPS L2 P(Y)相关通道"],
            ["等效矩形预相关带宽", "20.5 MHz", "报告附件2估算值"],
            ["N_LIM", "1", "硬限幅型模型，可用简化公式"],
            ["基线PDC", "0.0793", "既有脉冲干扰最坏背景"],
            ["基线R_I", "0", "基线弱脉冲项"],
            ["基线I_0,WB/N_0", "0.3925", "既有连续宽带干扰背景"],
            ["允许新增劣化", "0.2 dB", "M.2030表2/RS.2537表9"],
        ],
        widths=[5.0, 4.2, 7.1],
    )

    add_heading(doc, "4  第一步：将峰值功率换算为dBW并计算EIRP", 1)
    add_text(doc, "SAR5峰值功率为6120 W。功率从瓦换算为dBW时使用10log10(P/W)：")
    add_equation(doc, ["P_tx,pk = 10 log₁₀(6120)", "P_tx,pk ≈ 37.87 dBW ≈ 37.9 dBW"])
    add_text(doc, "加上发射天线峰值增益36.6 dBi，得到峰值等效全向辐射功率：")
    add_equation(doc, ["EIRP_pk = P_tx,pk + G_tx", "EIRP_pk = 37.9 + 36.6 = 74.5 dBW"])
    add_text(doc, "该结果与RS.2537-0表1中SAR5的峰值EIRP一致。")

    add_heading(doc, "5  第二步：计算SBAS接收机输入端峰值干扰功率", 1)
    add_text(doc, "报告的最坏情况假设是：RNSS接收机位于SAR主波束内，采用峰值发射功率和峰值发射增益，不计RNSS滤波。峰值链路预算为：")
    add_equation(doc, "P_RFI,pk = P_tx,pk + G_tx + G_rx − L_fs − L_pol − L_other")
    add_text(doc, "本例未单列其他附加损耗，因此取L_other = 0。代入报告表11参数：")
    add_equation(doc, [
        "P_RFI,pk = 37.9 + 36.6 + (−2.0) − 150.5 − 1.46",
        "P_RFI,pk ≈ −79.46 dBW",
        "报告表11舍入结果：P_RFI,pk = −79.48 dBW",
    ])

    add_heading(doc, "6  第三步：判断是否超过接收机生存电平", 1)
    add_text(doc, "SBAS接收机生存电平为−10 dBW。峰值干扰为−79.48 dBW，二者之间的生存裕量为：")
    add_equation(doc, [
        "M_survival = P_survival − P_RFI,pk",
        "M_survival = −10 − (−79.48) = 69.48 dB",
    ])
    add_text(doc, "因此，SAR5峰值干扰比接收机损伤门限低约69.5 dB，不会造成接收机前端烧毁或永久损伤。", bold=True, color="008000")
    add_text(doc, "这里应注意：dBW为对数功率，数值越接近0表示功率越大。因此−79.48 dBW远小于−10 dBW。")

    add_heading(doc, "7  第四步：判断是否造成接收机饱和", 1)
    add_text(doc, "SBAS接收机输入饱和电平为−135 dBW。峰值干扰相对饱和电平的超量为：")
    add_equation(doc, [
        "ΔP_sat = P_RFI,pk − P_sat",
        "ΔP_sat = −79.48 − (−135) = 55.52 dB",
    ])
    add_text(doc, "SAR脉冲到来时，峰值干扰高出接收机饱和电平约55.5 dB，因此接收机将进入强饱和状态。", bold=True, color="C00000")
    add_text(doc, "RS.2537-0还定义了进入性能劣化计算的筛查线：")
    add_equation(doc, ["P_screen = P_sat − 15 dB", "P_screen = −135 − 15 = −150 dBW"])
    add_text(doc, "实际峰值−79.48 dBW显著高于−150 dBW，所以必须继续进行PDC和性能劣化评估。结论可概括为：不会烧毁，但会强烈饱和。")

    add_heading(doc, "8  第五步：计算SAR信号与SBAS滤波器的频谱重叠", 1)
    add_heading(doc, "8.1  SBAS预相关滤波器频率范围", 2)
    add_text(doc, "SBAS等效矩形滤波器中心频率为1227.6 MHz，带宽为20.5 MHz。")
    add_equation(doc, [
        "f_rx,L = 1227.6 − 20.5/2 = 1217.35 MHz",
        "f_rx,H = 1227.6 + 20.5/2 = 1237.85 MHz",
        "SBAS等效通带：1217.35–1237.85 MHz",
    ])

    add_heading(doc, "8.2  SAR5调频脉冲频率范围", 2)
    add_text(doc, "SAR5中心频率为1236.5 MHz，Chirp带宽为14 MHz。")
    add_equation(doc, [
        "f_SAR,L = 1236.5 − 14/2 = 1229.5 MHz",
        "f_SAR,H = 1236.5 + 14/2 = 1243.5 MHz",
        "SAR频谱范围：1229.5–1243.5 MHz",
    ])

    add_heading(doc, "8.3  计算交叠带宽与重叠比", 2)
    add_text(doc, "两个频带的交集为1229.5–1237.85 MHz，因此：")
    add_equation(doc, [
        "Δf = 1237.85 − 1229.5 = 8.35 MHz",
        "η = Δf / B_chirp = 8.35 / 14 = 0.59643 ≈ 0.5964",
    ])
    add_text(doc, "因此，14 MHz线性调频脉冲中约59.64%的扫频时间落入SBAS等效接收通带。报告表13给出的重叠比正是0.5964。")

    add_heading(doc, "9  第六步：计算有效脉冲宽度", 1)
    add_text(doc, "对于线性调频信号，报告将落入RNSS预相关通带的那一部分脉冲时间定义为有效脉宽：")
    add_equation(doc, "PW_eff = PW × (Δf / B_chirp)")
    add_equation(doc, [
        "PW_eff = 37.63 μs × 0.5964",
        "PW_eff ≈ 22.445 μs",
    ])
    add_text(doc, "这意味着，原始脉冲虽然持续37.63 μs，但只有约22.445 μs对应的瞬时频率落入SBAS通带，不能直接用完整脉宽计算有效PDC。")

    add_heading(doc, "10  第七步：加入接收机过载恢复时间", 1)
    add_text(doc, "接收机在干扰脉冲结束后并不能立即恢复。SBAS恢复时间τ_r = 1 μs，因此每个SAR脉冲造成的总信号损失时间为：")
    add_equation(doc, [
        "PW_loss = PW_eff + τ_r",
        "PW_loss = 22.445 + 1 = 23.445 μs",
    ])

    add_heading(doc, "11  第八步：计算新增SAR5脉冲源的PDC", 1)
    add_text(doc, "RS.2537-0公式（5）为：")
    add_equation(doc, "PDC_Y = (PW_eff + τ_r) × PRF")
    add_text(doc, "将23.445 μs换算为秒并代入PRF = 1860 Hz：")
    add_equation(doc, [
        "PDC_Y = 23.445 × 10⁻⁶ × 1860",
        "PDC_Y = 0.0436077",
        "PDC_Y = 4.36077% ≈ 4.361%",
    ])
    add_text(doc, "该结果与报告表13中的4.361%一致。")

    add_heading(doc, "11.1  PDC中脉冲与恢复时间的贡献", 2)
    add_equation(doc, [
        "PDC_pulse = 22.445 × 10⁻⁶ × 1860 = 0.0417477 = 4.17477%",
        "PDC_recovery = 1 × 10⁻⁶ × 1860 = 0.00186 = 0.186%",
        "PDC_Y = 4.17477% + 0.186% = 4.36077%",
    ])
    add_text(doc, "恢复时间虽然只有1 μs，但由于每秒重复1860次，仍会使PDC增加0.186个百分点。")

    add_heading(doc, "11.2  与SAR原始发射占空比的区别", 2)
    add_equation(doc, [
        "DC_tx = PW × PRF",
        "DC_tx = 37.63 × 10⁻⁶ × 1860 ≈ 0.06999 ≈ 7.0%",
    ])
    add_table(
        doc,
        ["占空比定义", "结果", "含义"],
        [
            ["SAR原始发射占空比", "约7.0%", "完整脉宽×PRF"],
            ["通带内脉冲占空比", "约4.175%", "仅计频谱重叠部分"],
            ["有效丢失PDC", "约4.361%", "通带内脉宽+恢复时间"],
        ],
        widths=[5.0, 3.5, 7.6],
    )
    add_text(doc, "因此，RS.2537-0并不是把SAR的7%发射占空比直接代入，而是计算“接收机实际失去有效导航信号的占空比”。")

    add_heading(doc, "12  第九步：与原有基线脉冲干扰聚合", 1)
    add_text(doc, "SBAS接收机所处环境中还存在地面监视雷达等既有脉冲源。报告表9给出的基线PDC为0.0793。对于时间上可能重叠的不同强脉冲源，不能简单相加，而应按未受干扰时间的补集相乘：")
    add_equation(doc, [
        "1 − PDC_base+Y = (1 − PDC_base)(1 − PDC_Y)",
        "PDC_base+Y = 1 − (1 − 0.0793)(1 − 0.04361)",
        "PDC_base+Y = 1 − 0.9207 × 0.95639",
        "PDC_base+Y ≈ 0.119452 = 11.945%",
    ])
    add_text(doc, "若直接相加会得到12.291%，而补集乘积结果为11.945%。差值反映了基线脉冲和SAR脉冲在时间上重叠的可能性。对于低于门限的弱脉冲，参数R_I则按功率密度比直接求和。")

    add_heading(doc, "13  第十步：选择正确的接收机劣化公式", 1)
    add_text(doc, "SBAS接收机的N_LIM = 1，且本例将SAR5视为超过饱和门限的强脉冲源，因此R_Y≈0；表9中的基线R_I也为0。此时可采用报告公式（8）的硬限幅简化形式：")
    add_equation(doc, "N₀,EFF+Y / N₀,EFF = 1 / (1 − PDC_Y)²")
    add_text(doc, "这里N₀,EFF表示加入SAR5以前的有效噪声密度，N₀,EFF+Y表示加入SAR5以后。虽然表9还给出了基线PDC和连续干扰比I₀,WB/N₀，但在N_LIM = 1、R_I = R_Y = 0的增量劣化比中，公共基线项相互约掉，最终只与新增源PDC_Y有关。")
    add_text(doc, "这并不表示原有基线干扰不存在，而是表示0.2 dB允许值针对的是“在既定基线环境上，新增SAR5造成的额外劣化”。", bold=True, color="9C5700")

    add_heading(doc, "14  第十一步：计算线性有效噪声劣化比", 1)
    add_equation(doc, [
        "D_linear = 1 / (1 − 0.04361)²",
        "D_linear = 1 / 0.95639²",
        "D_linear ≈ 1.09328",
    ])
    add_text(doc, "加入SAR5后，有效噪声与干扰密度相对于基线增加：")
    add_equation(doc, "(1.09328 − 1) × 100% = 9.328%")
    add_text(doc, "即有效噪声密度约增加9.33%。")

    add_heading(doc, "15  第十二步：转换为dB劣化", 1)
    add_text(doc, "报告以有效噪声密度比的dB值作为新增脉冲干扰劣化量：")
    add_equation(doc, [
        "D_dB = 10 log₁₀(N₀,EFF+Y / N₀,EFF)",
        "D_dB = 10 log₁₀[1 / (1 − PDC_Y)²]",
        "D_dB = −20 log₁₀(1 − PDC_Y)",
    ])
    add_equation(doc, [
        "D_dB = −20 log₁₀(1 − 0.04361)",
        "D_dB = −20 log₁₀(0.95639)",
        "D_dB ≈ 0.387 dB",
    ])
    add_text(doc, "从性能角度看，也可以表述为：接收机有效C/N₀相对于基线下降约0.387 dB。该数值与RS.2537-0附件2第3.1.2节完全一致。")

    add_heading(doc, "16  第十三步：与允许劣化值比较", 1)
    add_equation(doc, [
        "D_allow = 0.2 dB",
        "D_actual = 0.387 dB",
        "ΔD = 0.387 − 0.2 = 0.187 dB",
    ])
    add_text(doc, "由于0.387 dB > 0.2 dB，本工况在RS.2537-0/M.2030初步评估中判定为超限。", bold=True, color="C00000")
    add_text(doc, "严格结论应写成：SAR5采用1236.5 MHz中心频率、14 MHz Chirp、37.63 μs脉宽和1860 Hz PRF时，对SBAS地面参考接收机造成的新增脉冲干扰劣化超过0.2 dB允许值，应转入更详细分析。")
    add_text(doc, "初评超限不能直接等同于“两系统绝对不能共存”。更详细的分析应引入实际轨道、时变方向图、真实接收滤波器、工作模式、照射持续时间和多星聚合。", bold=True, color="9C5700")

    add_heading(doc, "17  为什么该算例是保守的", 1)
    add_text(doc, "附件2为展示解析方法，假设在至少数分钟的某一时间段内，SAR5接收峰值持续高于SBAS饱和电平，使接收机在每个脉冲期间都处于饱和。真实场景中，接收功率应是时间函数：")
    add_equation(doc, "P_RFI(t) = P_tx + G_tx(t) + G_rx(t) − L_fs(t) − L_pol − L_misc")
    add_bullet(doc, "星地斜距随过境变化；")
    add_bullet(doc, "SAR波束仅在特定成像条带和观察角下照射接收机；")
    add_bullet(doc, "发射天线增益随主瓣、旁瓣和后瓣位置变化；")
    add_bullet(doc, "RNSS天线增益随卫星仰角和入射方向变化；")
    add_bullet(doc, "不同SAR模式、中心频率和带宽会改变频谱重叠；")
    add_bullet(doc, "真实接收滤波器不是理想矩形，带外阻塞与滤波裙边需另行计入。")
    add_text(doc, "因此，正常的详细分析应通过轨道仿真、SAR辐射方向图和RNSS接收天线方向图求得P_RFI(t)，再统计饱和事件的持续时间、PDC时间历程和超限概率。")

    add_heading(doc, "18  同一案例中中心频率的影响", 1)
    add_table(
        doc,
        ["SAR5中心频率", "14 MHz频谱范围", "与SBAS通带重叠", "初步结果"],
        [
            ["1236.5 MHz", "1229.5–1243.5 MHz", "8.35 MHz", "PDC=4.361%，劣化0.387 dB，超限"],
            ["1257.5 MHz", "1250.5–1264.5 MHz", "0 MHz", "理想矩形模型中PDC=0"],
            ["1278.5 MHz", "1271.5–1285.5 MHz", "0 MHz", "理想矩形模型中PDC=0"],
        ],
        widths=[3.6, 4.5, 3.6, 5.6],
    )
    add_text(doc, "报告由此认为，在其矩形滤波器近似下，SAR5采用1257.5或1278.5 MHz时，不会通过SBAS预相关通带形成该类脉冲劣化。")
    add_text(doc, "但PDC=0只表示理想通带内无重叠，并不代表真实设备绝对无干扰。详细分析仍需考虑接收滤波器裙边、SAR带外谱、功放非线性谱再生、前端阻塞和互调。")

    add_heading(doc, "19  工程反算：满足0.2 dB时允许的最大PDC", 1)
    add_text(doc, "以下是基于报告公式的工程反算，不是RS.2537-0直接列出的数值。要求：")
    add_equation(doc, "−20 log₁₀(1 − PDC_max) ≤ 0.2")
    add_equation(doc, [
        "1 − PDC_max = 10^(−0.2/20)",
        "PDC_max = 1 − 10^(−0.01)",
        "PDC_max ≈ 0.02276 = 2.276%",
    ])
    add_text(doc, "本例实际PDC为4.361%，因此需要降低到当前值的：")
    add_equation(doc, [
        "2.276% / 4.361% ≈ 0.522",
        "即PDC需降低约47.8%",
    ])

    add_heading(doc, "20  工程反算：允许的最大频谱重叠比", 1)
    add_text(doc, "保持脉宽37.63 μs、PRF 1860 Hz和恢复时间1 μs不变，由PDC = (PW·η + τ_r)PRF可得：")
    add_equation(doc, [
        "η_max = [PDC_max/PRF − τ_r] / PW",
        "η_max = [0.0227628/1860 − 1×10⁻⁶] / (37.63×10⁻⁶)",
        "η_max ≈ 0.2986",
    ])
    add_text(doc, "对应14 MHz Chirp允许的最大重叠带宽：")
    add_equation(doc, [
        "Δf_max = 14 × 0.2986",
        "Δf_max ≈ 4.18 MHz",
    ])
    add_text(doc, "本例实际重叠为8.35 MHz。因此在其他条件不变的理想矩形滤波器模型下，需要将重叠宽度降至约4.18 MHz以下，才可能满足0.2 dB初评要求。")

    add_heading(doc, "21  完整结果汇总", 1)
    add_table(
        doc,
        ["计算环节", "计算结果", "判定或含义"],
        [
            ["SAR5峰值EIRP", "74.5 dBW", "峰值发射功率+天线增益"],
            ["SBAS输入峰值干扰", "−79.48 dBW", "报告表11结果"],
            ["SBAS生存电平", "−10 dBW", "不会损坏，生存裕量69.48 dB"],
            ["SBAS饱和电平", "−135 dBW", "峰值高出55.52 dB，强饱和"],
            ["频谱重叠", "8.35 MHz", "重叠比0.5964"],
            ["有效脉宽", "22.445 μs", "仅计通带内Chirp时间"],
            ["单脉冲总损失时间", "23.445 μs", "含1 μs恢复时间"],
            ["新增SAR5 PDC", "4.361%", "报告表13一致"],
            ["加入基线后的总PDC", "11.945%", "补集乘积聚合"],
            ["有效噪声密度增加", "约9.33%", "线性劣化比1.09328"],
            ["新增dB劣化", "0.387 dB", "相当于C/N₀下降约0.387 dB"],
            ["允许值", "0.2 dB", "SBAS新增脉冲源预算"],
            ["初评结论", "超限0.187 dB", "需要详细动态分析"],
        ],
        widths=[5.2, 4.2, 7.0],
    )

    add_heading(doc, "22  方法的核心理解", 1)
    add_numbered(doc, "会不会烧毁接收机？本例不会，峰值相对生存电平仍有约69.5 dB裕量。")
    add_numbered(doc, "会不会让接收机饱和？会，峰值高出饱和电平约55.5 dB。")
    add_numbered(doc, "每个脉冲实际让RNSS丢失多长时间？不是完整37.63 μs，而是频谱重叠后的22.445 μs，再加1 μs恢复时间。")
    add_numbered(doc, "这些丢失时间累积后是否超过任务预算？有效PDC为4.361%，对应0.387 dB新增劣化，高于0.2 dB。")
    add_text(doc, "因此，RS.2537-0方法的实质不是简单比较“SAR峰值功率”和“RNSS干扰门限”，而是把前端生存、前端饱和、频谱重叠、接收机恢复、时间占用和既有干扰背景串成一条完整判定链。", bold=True, color="17365D")
    add_text(doc, "当存在多个星载有源传感器可能同时照射同一RNSS接收机时，还必须考虑聚合影响。报告指出，可通过EESS（有源）系统运营方之间的运行协调降低潜在的聚合干扰。")

    add_heading(doc, "附录A  关键公式速查", 1)
    add_table(
        doc,
        ["用途", "公式"],
        [
            ["峰值干扰链路", "P_RFI,pk = P_tx,pk + G_tx + G_rx − L_fs − L_pol − L_other"],
            ["有效脉宽", "PW_eff = PW × (Δf/B_chirp)"],
            ["强脉冲PDC", "PDC_j = (PW_eff + τ_r) × PRF_j"],
            ["弱脉冲占空比", "dc_k = PW_eff × PRF_k"],
            ["多强源PDC聚合", "PDC_total = 1 − ∏(1 − PDC_j)"],
            ["多弱源R_I聚合", "R_I,total = ΣR_I,j"],
            ["硬限幅新增劣化", "N₀,EFF+Y/N₀,EFF = 1/(1 − PDC_Y)²"],
            ["dB劣化", "D_dB = −20log₁₀(1 − PDC_Y)"],
        ],
        widths=[5.2, 11.2],
    )

    add_heading(doc, "附录B  原报告定位", 1)
    add_table(
        doc,
        ["内容", "RS.2537-0位置"],
        [
            ["六类SAR技术参数", "附件1表1，报告印刷第5–6页"],
            ["分析方法：损坏/饱和", "附件1第4.1节，报告印刷第15页"],
            ["PDC、R_I和有效脉宽", "附件1第4.2节，报告印刷第16–18页"],
            ["SBAS基线参数与0.2 dB允许值", "附件1表9，报告印刷第19页"],
            ["SBAS峰值干扰链路", "附件2表11，报告印刷第22页"],
            ["SAR5—SBAS PDC计算", "附件2表13/14，报告印刷第24页"],
            ["0.387 dB和0.349 dB结果", "附件2第3.1.2节，报告印刷第25页"],
        ],
        widths=[7.8, 8.6],
    )

    add_heading(doc, "参考文献", 1)
    add_text(doc, "[1] Report ITU-R RS.2537-0, Representative system characteristics and examples of evaluating interference into receiving earth stations in the radionavigation-satellite service (space-to-Earth) from spaceborne synthetic aperture radar sensors in the Earth exploration-satellite (active) service in the 1 215–1 300 MHz band, 2023.", indent=False, size=9.5)
    add_text(doc, "[2] Recommendation ITU-R M.2030-0, Evaluation method for pulsed interference from relevant radio sources other than in the radionavigation-satellite service to RNSS systems and networks, 2012.", indent=False, size=9.5)

    # Keep tables together where possible
    for table in doc.tables:
        for row in table.rows:
            tr_pr = row._tr.get_or_add_trPr()
            cant_split = OxmlElement("w:cantSplit")
            tr_pr.append(cant_split)

    doc.core_properties.title = "ITU-R RS.2537-0方法详细演算"
    doc.core_properties.subject = "SAR5对GPS SBAS地面参考接收机脉冲干扰初步评估"
    doc.core_properties.author = "OpenAI"
    doc.core_properties.keywords = "ITU-R, RS.2537-0, SAR5, SBAS, PDC, 脉冲干扰, RNSS"
    return doc


if __name__ == "__main__":
    document = create_document()
    document.save(OUT_PATH)
    print(OUT_PATH)
