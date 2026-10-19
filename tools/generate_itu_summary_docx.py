from __future__ import annotations

from pathlib import Path
from typing import Iterable

import matplotlib.pyplot as plt
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT_DIR = Path("generated")
CHART_DIR = OUTPUT_DIR / "charts_m2220_m2305"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
CHART_DIR.mkdir(parents=True, exist_ok=True)
DOCX_PATH = OUTPUT_DIR / "ITU-R_M.2220-1与M.2305-1详细解读.docx"

BODY_FONT = "Noto Sans CJK SC"
HEADING_FONT = "Noto Serif CJK SC"
MATH_FONT = "Cambria Math"
ACCENT = "1F4E78"
ACCENT_2 = "4472C4"
ACCENT_LIGHT = "D9EAF7"
ACCENT_VERY_LIGHT = "EEF5FA"
GRAY = "666666"
LIGHT_GRAY = "F2F2F2"
DARK = "1F1F1F"
ORANGE_LIGHT = "FCE4D6"
GREEN_LIGHT = "E2F0D9"

URL_M2220 = "https://www.itu.int/dms_pub/itu-r/opb/rep/R-REP-M.2220-1-2021-PDF-E.pdf"
URL_M2305 = "https://www.itu.int/dms_pub/itu-r/opb/rep/R-REP-M.2305-1-2026-PDF-E.pdf"
URL_M2030 = "https://www.itu.int/rec/R-REC-M.2030/en"
URL_RS2537 = "https://www.itu.int/pub/R-REP-RS.2537-2023"


def set_run_font(run, font_name: str = BODY_FONT, size: float | None = None,
                 bold: bool | None = None, color: str | None = None,
                 italic: bool | None = None) -> None:
    run.font.name = font_name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font_name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_spacing(paragraph, before: float = 0, after: float = 6,
                          line: float = 1.42) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def set_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r0 = paragraph.add_run("第 ")
    set_run_font(r0, size=9, color=GRAY)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r = paragraph.add_run()
    r._r.extend([fld_begin, instr, fld_end])
    r1 = paragraph.add_run(" 页")
    set_run_font(r1, size=9, color=GRAY)


def add_toc(paragraph) -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "打开文档后右击目录并选择“更新域—更新整个目录”。"
    fld_sep.append(text)
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, fld_end])


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([color, underline])
    new_run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    r = p.add_run(text)
    set_run_font(r, HEADING_FONT, bold=True, color=ACCENT)
    if level == 1:
        r.font.size = Pt(16)
        p.paragraph_format.space_before = Pt(15)
        p.paragraph_format.space_after = Pt(8)
    elif level == 2:
        r.font.size = Pt(13)
        p.paragraph_format.space_before = Pt(11)
        p.paragraph_format.space_after = Pt(5)
    else:
        r.font.size = Pt(11.5)
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True


def add_body(doc: Document, text: str, *, first_indent: bool = True,
             bold_prefix: str | None = None, keep_with_next: bool = False) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0.74) if first_indent else Cm(0)
    p.paragraph_format.keep_with_next = keep_with_next
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    set_paragraph_spacing(p, after=6, line=1.45)


def add_bullet(doc: Document, text: str, level: int = 0) -> None:
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.left_indent = Cm(0.75 + 0.55 * level)
    p.paragraph_format.first_line_indent = Cm(-0.45)
    r = p.add_run(text)
    set_run_font(r)
    set_paragraph_spacing(p, after=3, line=1.32)


def add_numbered(doc: Document, text: str, level: int = 0) -> None:
    style = "List Number" if level == 0 else "List Number 2"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.left_indent = Cm(0.75 + 0.55 * level)
    p.paragraph_format.first_line_indent = Cm(-0.45)
    r = p.add_run(text)
    set_run_font(r)
    set_paragraph_spacing(p, after=3, line=1.32)


def add_formula(doc: Document, formula: str, explanation: str | None = None) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, ACCENT_VERY_LIGHT)
    set_cell_margins(cell, top=120, start=170, bottom=120, end=170)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(formula)
    set_run_font(r, MATH_FONT, size=11.5, color=DARK)
    if explanation:
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run(explanation)
        set_run_font(r2, size=9, color=GRAY)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)


def add_callout(doc: Document, title: str, body: str,
                fill: str = ACCENT_LIGHT) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=130, start=170, bottom=130, end=170)
    p = cell.paragraphs[0]
    r = p.add_run(title)
    set_run_font(r, bold=True, color=ACCENT)
    p2 = cell.add_paragraph()
    r2 = p2.add_run(body)
    set_run_font(r2)
    set_paragraph_spacing(p2, after=0, line=1.38)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)


def add_source_note(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.right_indent = Cm(0.5)
    r1 = p.add_run("原文依据：")
    set_run_font(r1, size=8.5, bold=True, color=GRAY)
    r2 = p.add_run(text)
    set_run_font(r2, size=8.5, color=GRAY)
    set_paragraph_spacing(p, after=5, line=1.2)


def add_table(doc: Document, headers: list[str], rows: Iterable[Iterable[str]],
              widths_cm: list[float] | None = None, font_size: float = 8.8,
              first_col_center: bool = True) -> None:
    data = list(rows)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, ACCENT)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_run_font(r, size=font_size, bold=True, color="FFFFFF")
        if widths_cm:
            cell.width = Cm(widths_cm[i])
    for ridx, row_values in enumerate(data):
        row = table.add_row()
        set_cant_split(row)
        for i, value in enumerate(row_values):
            cell = row.cells[i]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ridx % 2 == 1:
                set_cell_shading(cell, ACCENT_VERY_LIGHT)
            p = cell.paragraphs[0]
            if i == 0 and first_col_center:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            set_run_font(r, size=font_size)
            if widths_cm:
                cell.width = Cm(widths_cm[i])
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)


def add_figure(doc: Document, image_path: Path, caption: str,
               width_cm: float = 15.8) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(image_path), width=Cm(width_cm))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(caption)
    set_run_font(r, size=9, color=GRAY, italic=True)
    set_paragraph_spacing(c, after=7, line=1.1)


def configure_document(doc: Document) -> None:
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.05)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.25)
    sec.right_margin = Cm(2.05)
    sec.header_distance = Cm(0.9)
    sec.footer_distance = Cm(0.9)

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.45
    normal.paragraph_format.space_after = Pt(6)

    for level, size in ((1, 16), (2, 13), (3, 11.5)):
        st = doc.styles[f"Heading {level}"]
        st.font.name = HEADING_FONT
        st._element.rPr.rFonts.set(qn("w:eastAsia"), HEADING_FONT)
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(ACCENT)

    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run("ITU-R M.2220-1与M.2305-1详细解读")
    set_run_font(hr, size=8.5, color=GRAY)
    add_page_number(sec.footer.paragraphs[0])


def generate_charts() -> dict[str, Path]:
    paths: dict[str, Path] = {}

    # M.2220 representative aggregate parameters.
    labels = ["US high-alt", "Europe high-alt", "DCA low-alt"]
    pdc = [0.6527, 0.6031, 0.3123]
    ri = [0.9628, 1.5218, 0.4524]
    x = list(range(len(labels)))
    width = 0.34
    fig = plt.figure(figsize=(8.2, 4.5))
    ax = fig.add_subplot(111)
    ax.bar([v - width / 2 for v in x], pdc, width, label="Aggregate PDC")
    ax.bar([v + width / 2 for v in x], ri, width, label="Aggregate RI")
    ax.set_xticks(x, labels)
    ax.set_ylabel("Dimensionless parameter")
    ax.set_title("Representative M.2220 airborne aggregate parameters")
    ax.grid(axis="y", alpha=0.3)
    ax.legend()
    fig.tight_layout()
    path = CHART_DIR / "m2220_airborne_parameters.png"
    fig.savefig(path, dpi=180)
    plt.close(fig)
    paths["m2220"] = path

    lat = [40, 50, 60, 70, 80]
    study_a = {
        "A1 SAR3+SAR6": [0, 0, 0, 0, 41],
        "A2 SAR3+Scat2": [91, 82, 87, 139, 409],
        "A3 SAR6+Scat2": [0, 0, 0, 6, 1186],
    }
    fig = plt.figure(figsize=(8.2, 4.6))
    ax = fig.add_subplot(111)
    for name, values in study_a.items():
        ax.plot(lat, values, marker="o", label=name)
    ax.set_yscale("symlog", linthresh=1)
    ax.set_xlabel("Receiver latitude (deg)")
    ax.set_ylabel("Event occurrences / scans")
    ax.set_title("M.2305 Study A: strong-strong event occurrences in 120 days")
    ax.grid(True, alpha=0.3)
    ax.legend(fontsize=8)
    fig.tight_layout()
    path = CHART_DIR / "m2305_study_a_events.png"
    fig.savefig(path, dpi=180)
    plt.close(fig)
    paths["study_a"] = path

    study_b = {
        "B4 widebeam orthogonal": [0, 0, 0, 32, 194],
        "B5 widebeam co-planar": [44, 55, 68, 116, 189],
        "B6 line-of-sight": [0, 25, 185, 296, 459],
    }
    fig = plt.figure(figsize=(8.2, 4.6))
    ax = fig.add_subplot(111)
    for name, values in study_b.items():
        ax.plot(lat, values, marker="o", label=name)
    ax.set_yscale("symlog", linthresh=1)
    ax.set_xlabel("Receiver latitude (deg)")
    ax.set_ylabel("Event occurrences")
    ax.set_title("M.2305 Study B: potential future aggregate events in 120 days")
    ax.grid(True, alpha=0.3)
    ax.legend(fontsize=8)
    fig.tight_layout()
    path = CHART_DIR / "m2305_study_b_events.png"
    fig.savefig(path, dpi=180)
    plt.close(fig)
    paths["study_b"] = path

    # Annex 2 degradation example.
    fig = plt.figure(figsize=(7.5, 4.2))
    ax = fig.add_subplot(111)
    names = ["SAR1", "SARA", "Simultaneous"]
    values = [0.019, 0.196, 0.215]
    ax.bar(names, values)
    ax.axhline(0.2, linestyle="--", label="Allowable 0.2 dB")
    ax.set_ylabel("Additional degradation (dB)")
    ax.set_title("M.2305 Annex 2: individual pass, aggregate fail")
    ax.grid(axis="y", alpha=0.3)
    ax.legend()
    fig.tight_layout()
    path = CHART_DIR / "m2305_annex2_degradation.png"
    fig.savefig(path, dpi=180)
    plt.close(fig)
    paths["annex2"] = path

    return paths


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(58)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("ITU-R M.2220-1与M.2305-1")
    set_run_font(r, HEADING_FONT, size=25, bold=True, color=ACCENT)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("脉冲干扰参数计算与多星聚合事件分析\n详细解读")
    set_run_font(r2, HEADING_FONT, size=21, bold=True, color=ACCENT)

    line = doc.add_paragraph()
    line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = line.add_run("━━━━━━━━━━━━━━━━━━━━━━━━━━━━")
    set_run_font(rr, size=14, color=ACCENT_2)

    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_before = Pt(17)
    r3 = p3.add_run(
        "面向1215–1300 MHz星载有源传感器与RNSS接收机共存研究\n"
        "解释两份报告做了什么、怎样做、得到什么结果，以及如何用于项目"
    )
    set_run_font(r3, size=12.5, color=DARK)

    table = doc.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info = [
        ("报告一", "Report ITU-R M.2220-1（11/2021）"),
        ("报告二", "Report ITU-R M.2305-1（11/2025，2026电子出版）"),
        ("方法主线", "源—路径—接收机参数化 → 强/弱脉冲聚合 → 轨道与多星同时照射统计"),
        ("文档性质", "基于ITU-R原文的中文技术解读；工程延伸内容已单独标注"),
    ]
    for i, (k, v) in enumerate(info):
        set_cell_shading(table.cell(i, 0), ACCENT_LIGHT)
        set_cell_margins(table.cell(i, 0), top=120, start=130, bottom=120, end=130)
        set_cell_margins(table.cell(i, 1), top=120, start=130, bottom=120, end=130)
        table.cell(i, 0).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        rk = table.cell(i, 0).paragraphs[0].add_run(k)
        set_run_font(rk, bold=True, color=ACCENT)
        rv = table.cell(i, 1).paragraphs[0].add_run(v)
        set_run_font(rv)

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_before = Pt(35)
    rn = note.add_run("整理日期：2026年10月19日")
    set_run_font(rn, size=10, color=GRAY)
    doc.add_page_break()


def add_front_matter(doc: Document) -> None:
    add_heading(doc, "阅读说明", 1)
    add_body(
        doc,
        "M.2220-1和M.2305-1均为ITU-R研究报告（Report），不是具有“做出建议”条款的建议书。"
        "前者解决“怎样把复杂脉冲环境转化为M.2030所需的PDC、RI等参数”，后者解决“多个星载有源传感器何时会同时照射同一RNSS接收机、这种事件发生多少次和持续多久”。"
        "它们互为补充，但都不能单独完成最终兼容性判定。",
    )
    add_callout(
        doc,
        "最简记忆",
        "M.2220-1是“参数计算引擎”：把发射机、传播、天线、滤波器和接收机门限转换为聚合PDC与RI。\n"
        "M.2305-1是“事件统计引擎”：用轨道、波束和时间仿真找出多星同时照射事件的次数与持续时间。\n"
        "M.2030才是“性能判定引擎”：把PDC、RI和连续干扰基线转换成RNSS性能劣化。",
        GREEN_LIGHT,
    )

    add_heading(doc, "目录", 1)
    toc = doc.add_paragraph()
    add_toc(toc)
    doc.add_page_break()

    add_heading(doc, "执行摘要", 1)
    add_table(
        doc,
        ["比较项", "M.2220-1", "M.2305-1"],
        [
            ("核心问题", "真实脉冲环境如何参数化？", "多个EESS传感器何时形成同时照射？"),
            ("主要输入", "发射机、位置、天线、波形、传播、接收机滤波与门限", "轨道、波束、天线方向图、接收机位置、强脉冲门限"),
            ("主要输出", "聚合PDC/PDCLIM、RI及连续干扰参数", "事件次数、持续时间、纬度分布和同时照射窗口"),
            ("主要工具", "源—路径—接收机模型、GREET、碰撞模型、解析计算", "STK轨道仿真、RIP记录、MATLAB后处理"),
            ("能否直接判定合格", "不能；输出送入M.2030", "不能；需用实际参数并接M.2220/M.2030"),
            ("主要价值", "建立可复算的干扰参数基线", "把静态最坏值扩展为动态概率和持续时间"),
        ],
        widths_cm=[3.0, 6.4, 6.4],
        font_size=8.6,
    )
    add_body(
        doc,
        "两份报告组合以后，工程分析不再只是计算一个“最坏峰值功率”，而是形成一条可追溯链条：先识别每个脉冲源在接收机端是强脉冲还是弱脉冲，再计算聚合PDC和RI；随后用轨道级仿真确定这些干扰状态在时间和空间上何时同时出现；最后将逐时刻参数输入M.2030，计算劣化和超限统计。",
    )


def add_m2220_section(doc: Document, charts: dict[str, Path]) -> None:
    add_heading(doc, "第一篇：Report ITU-R M.2220-1详细解读", 1)

    add_heading(doc, "1. 报告要解决什么问题", 2)
    add_body(
        doc,
        "M.2220-1面对的不是一个单独雷达，而是由地面、机载和星载脉冲源共同构成的复杂环境。"
        "它提出一种统一计算方法，把所有带内和邻带脉冲源在RNSS接收机端的作用，压缩成一对可供M.2030使用的聚合参数：强脉冲时间损失参数PDC（或PDCLIM/PDCB）和弱脉冲等效噪声参数RI。",
    )
    add_body(
        doc,
        "报告覆盖两类频段：1164–1215 MHz主要面向航空RNSS接收机及DME/TACAN、CNI、ATC/TCAS环境；1215–1300 MHz主要面向地面RNSS接收机、地面监视雷达和星载SAR。"
        "报告同时给出高空、低空和地面接收机的计算实例，并用附件A说明GREET工具如何建立固定脉冲发射机数据库和传播环境。",
    )
    add_source_note(doc, "M.2220-1，§1、§1.3，PDF第4–6页；目录见PDF第3页。")

    add_heading(doc, "2. 它在整个ITU-R方法体系中的位置", 2)
    add_table(
        doc,
        ["层级", "文件/模型", "作用"],
        [
            ("环境建模", "M.2220-1", "把多个实际脉冲源转换成PDC、RI等模型输入"),
            ("性能计算", "M.2030", "根据PDC、RI、连续干扰和接收机NLIM计算有效噪声与新增劣化"),
            ("SAR算例", "RS.2537-0", "把M.2030/M.2220用于1215–1300 MHz星载SAR单源初评"),
            ("动态聚合", "M.2305-1", "判断多个EESS传感器同时照射事件的发生时间、次数和持续时间"),
        ],
        widths_cm=[2.8, 3.6, 9.3],
    )
    add_callout(
        doc,
        "容易混淆的地方",
        "M.2220-1不直接规定“允许劣化0.1 dB或0.2 dB”。它的任务是生成输入参数。"
        "保护限值和劣化公式来自M.2030及具体RNSS接收机建议书。",
        ORANGE_LIGHT,
    )

    add_heading(doc, "3. 基本框架：源—路径—接收机", 2)
    add_body(doc, "报告把分析拆成三个互相关联的模块：")
    add_numbered(doc, "源模型：发射机位置、工作频率、峰值功率、发射谱、脉宽、PRF、天线方向图和扫描方式。")
    add_numbered(doc, "路径模型：发射机与RNSS接收机的距离、视距条件、传播损耗、到达角及接收天线方向增益。")
    add_numbered(doc, "接收机模型：天线、RF/IF滤波器、脉冲熄减或饱和方式、饱和/熄减门限、过载恢复时间、AGC和ADC特性。")
    add_body(
        doc,
        "这种拆分的价值在于：同一个发射机并不天然属于“强脉冲”或“弱脉冲”。"
        "它经过传播损耗、天线增益和滤波衰减后，只有在特定几何位置和特定接收机上才会超过熄减或饱和门限。",
    )
    add_source_note(doc, "M.2220-1，§2.1，PDF第6页。")

    add_heading(doc, "4. 报告建立的四类运行场景", 2)
    add_table(
        doc,
        ["场景", "主要受扰对象", "主要脉冲源", "分析特点"],
        [
            ("高空航空场景", "航路飞行RNSS接收机", "地面DME/TACAN、CNI、ATC，机载DME/TCAS", "可见发射机多，聚合PDC高，重点关注热点区域"),
            ("低空进近场景", "终端区/进近RNSS接收机", "附近机载与地面脉冲源", "地面可见源减少，但机载源密集"),
            ("地面1164–1215 MHz", "机场地面、高精度/参考接收机", "有限数量DME/TACAN、ATC、CNI", "几何遮挡明显，接收机常按饱和模型处理"),
            ("地面1215–1300 MHz", "SBAS、高精度、航空导航等接收机", "地面监视雷达、星载SAR、连续业余业务", "需同时处理地面雷达和EESS脉冲源"),
        ],
        widths_cm=[3.0, 3.7, 4.5, 4.5],
        font_size=8.4,
    )
    add_source_note(doc, "M.2220-1，§2.1.1.1–2.1.1.4，PDF第6–8页。")

    add_heading(doc, "5. 两类接收机机理", 2)
    add_heading(doc, "5.1 脉冲熄减型接收机", 3)
    add_body(
        doc,
        "航空RNSS接收机可能使用快速数字脉冲熄减器。峰值超过门限的样本在相关处理前被置零，相当于损失一段有效积分时间；低于门限的脉冲则不被熄减，而表现为额外噪声。"
        "报告指出，在相同门限下，仅对脉冲幅度进行削顶通常比完全熄减带来约两倍的dB损失，因此接收机实现方式必须明确。",
    )
    add_heading(doc, "5.2 饱和型接收机", 3)
    add_body(
        doc,
        "许多地面RNSS接收机不设置快速熄减器。强脉冲会使RF前端、AGC或ADC短时饱和；在脉冲结束后，接收机还要经历过载恢复时间。"
        "此时实际信号损失时间是“高于门限的有效脉宽+恢复时间”，恢复时间越长，PDC越大。",
    )
    add_source_note(doc, "M.2220-1，§2.2.1，PDF第8–9页。")

    add_heading(doc, "6. 核心参数及其物理意义", 2)
    add_formula(
        doc,
        "PDC_j = (PW_j,eff + τ_r) × PRF_j",
        "单个强脉冲源造成的接收机失效时间比例；PW_eff为接收机端有效脉宽，τ_r为过载恢复时间。",
    )
    add_formula(
        doc,
        "PDC_agg = 1 − ∏_j (1 − PDC_j)",
        "多个强脉冲源按“未受干扰时间”的乘积聚合，避免简单相加重复计算时间重叠。",
    )
    add_formula(
        doc,
        "RI = Σ_i [P_i × dc_i / (N0 × BW)]",
        "多个低于门限的弱脉冲以平均功率密度相对于热噪声密度的比值求和。",
    )
    add_table(
        doc,
        ["参数", "表示什么", "受哪些因素控制", "常见错误"],
        [
            ("PDC/PDCB/PDCLIM", "强脉冲造成的净时间损失比例", "有效脉宽、PRF、恢复时间、脉冲碰撞", "把多个源直接相加，得到超过100%的占空比"),
            ("RI", "弱脉冲产生的等效噪声贡献", "峰值功率、占空比、接收机噪声和带宽", "只统计强脉冲，忽略门限以下的大量脉冲"),
            ("I0,WB/N0", "连续宽带干扰基线", "其他RNSS信号、连续发射和非必要发射", "把全部性能预算留给新脉冲源"),
            ("τr", "过载后恢复到可用状态所需时间", "AGC、ADC、滤波器、限幅器和实现方式", "把示例值1 μs当作所有接收机固定值"),
        ],
        widths_cm=[3.0, 4.3, 4.5, 4.1],
        font_size=8.3,
    )

    add_heading(doc, "7. 为什么不能把所有脉冲占空比直接相加", 2)
    add_body(
        doc,
        "在复杂航空环境中，不同系统的脉冲会随机重叠。简单相加既会重复计算重叠的强脉冲时间，也会把本来已被强脉冲熄减掉的弱脉冲功率再次计入RI。"
        "M.2220-1在附件B给出强脉冲碰撞模型，通过净PDC方法处理DME/TACAN、CNI、ATC等异质脉冲的重叠。",
    )
    add_body(
        doc,
        "对工程人员而言，其关键不是重现附件B的每一个概率推导，而是保持两个原则：强脉冲按时间并集计算；弱脉冲只在未被强脉冲熄减的时间内贡献等效噪声。",
    )

    add_heading(doc, "8. 各类脉冲源怎样建模", 2)
    add_table(
        doc,
        ["脉冲源", "报告中的模型重点", "为什么需要单独处理"],
        [
            ("DME/TACAN", "1 MHz信道、脉冲对、近似高斯脉冲；区分高于门限的脉宽和低于门限的高斯尾部", "数量多，是航空热点区域的主导聚合源"),
            ("CNI跳频系统", "51个载频、3 MHz间隔、TDMA与频率跳变；使用时间/空间密度因子", "单个终端出现概率和同时活动数决定聚合结果"),
            ("ATC/Mode S/TCAS", "地面与机载问答机、背景区域分布、前景强源和本机发射", "波形和业务状态复杂，脉冲序列不等价"),
            ("地面监视雷达", "主瓣/旁瓣EIRP、脉宽、PRF、距离和接收机饱和门限", "1215–1300 MHz地面接收机的主要强脉冲源"),
            ("星载SAR", "轨道、主瓣EIRP、线性调频带宽、脉宽、PRF及频谱重叠修正", "宽带chirp不一定全部落入接收机通带"),
            ("连续背景", "RNSS互系统/系统内信号及其他连续发射形成I0,WB/N0", "会减少新脉冲源可用的剩余干扰预算"),
        ],
        widths_cm=[2.8, 7.0, 6.1],
        font_size=8.3,
    )

    add_heading(doc, "9. 频谱重叠与扫描修正", 2)
    add_heading(doc, "9.1 线性调频脉冲的有效脉宽", 3)
    add_formula(
        doc,
        "PW_eff = PW × (BW_overlap / BW_chirp)",
        "当SAR调频带宽只有一部分落入RNSS预相关通带时，仅按重叠比例计入有效脉宽。",
    )
    add_body(
        doc,
        "这一步把“中心频率是否相同”升级为实际频谱耦合分析。详细评估应使用实测发射谱和真实接收滤波器，而不是只用矩形带宽。",
    )
    add_heading(doc, "9.2 快速扫描波束的有效PRF", 3)
    add_formula(
        doc,
        "PRF_eff = PRF × (τ_obs / T_c)",
        "某些极窄跟踪环路的半无代码接收机可用扫描照射时间比例修正；报告明确警告不能普遍套用。",
    )
    add_callout(
        doc,
        "重要限制",
        "扫描折减只在接收机跟踪环路极窄、其输出近似长期平均时成立。对于响应较快的现代接收机，"
        "瞬时C/N0下降不能被长期平均掩盖，需通过M.2496或实机试验确定。",
        ORANGE_LIGHT,
    )
    add_source_note(doc, "M.2220-1，§4.1.3–4.1.4，PDF第33–34页。")

    add_heading(doc, "10. GREET工具做了什么", 2)
    add_body(
        doc,
        "附件A介绍GPS RFI Environmental Evaluation Tool（GREET）。它把固定DME/TACAN台站数据库、RNSS接收机位置和高度、传播损耗、接收天线增益以及门限模型结合起来，逐站计算接收峰值功率。"
        "随后确定每个台站高于熄减门限的脉宽、净PDC和低于门限的RI，并可在经纬度网格上生成热点环境。",
    )
    add_numbered(doc, "导入固定发射机位置、频率、功率、天线和波形参数。")
    add_numbered(doc, "对每个接收机网格点计算距离、到达角、传播损耗和滤波衰减。")
    add_numbered(doc, "求每个脉冲在接收机输入端的峰值，以及高于门限的有效宽度。")
    add_numbered(doc, "按碰撞模型求净PDC，并对未熄减的脉冲尾部求RI。")
    add_numbered(doc, "输出区域热点、最坏网格点及可送入M.2030的聚合参数。")
    add_source_note(doc, "M.2220-1，附件A，PDF第38–42页。")

    add_heading(doc, "11. 航空接收机仿真结果怎样理解", 2)
    add_table(
        doc,
        ["代表场景", "聚合PDC", "聚合RI", "主要含义"],
        [
            ("美国高空热点", "0.6527", "0.9628", "DME/TACAN占主导，CNI和ATC继续增加时间损失及噪声"),
            ("欧洲高空热点", "0.6031", "1.5218", "强脉冲时间比例略低，但未熄减脉冲等效噪声更大"),
            ("华盛顿DCA低空进近", "0.3123", "0.4524", "低空可见地面源减少，聚合环境仍不可忽略"),
        ],
        widths_cm=[4.0, 2.7, 2.7, 6.5],
    )
    add_figure(
        doc,
        charts["m2220"],
        "图1  M.2220-1代表性航空场景的聚合PDC和RI（根据报告结果重绘；两者均为无量纲模型参数）",
    )
    add_body(
        doc,
        "PDC和RI不能直接横向比较大小，也不能直接解释成“检测概率下降百分比”。它们是不同机理的中间参数，必须和接收机NLIM、连续干扰基线及M.2030公式共同使用。",
    )

    add_heading(doc, "12. 地面接收机算例", 2)
    add_heading(doc, "12.1 SBAS地面参考接收机与两部监视雷达", 3)
    add_table(
        doc,
        ["项目", "雷达A", "雷达B"],
        [
            ("主瓣EIRP", "101 dBW", "85.6 dBW"),
            ("旁瓣EIRP", "61 dBW", "50.6 dBW"),
            ("脉宽", "2 μs", "51.2 μs × 2脉冲"),
            ("PRF", "358 Hz", "750脉冲对/s"),
            ("单源PDCLIM", "0.001074", "0.0783"),
        ],
        widths_cm=[5.0, 5.3, 5.3],
    )
    add_formula(
        doc,
        "PDCLIM,agg = 1 − (1 − 0.001074)(1 − 0.0783) = 0.0793",
        "由于强脉冲完全使接收机饱和，示例中RI近似为0。",
    )
    add_body(
        doc,
        "这个0.0793后来被M.2030和RS.2537用作SBAS接收机的基线PDC。它不是人为给定的经验值，而是从两部代表性监视雷达、接收机1 μs恢复时间和补集乘积聚合得到的。",
    )
    add_source_note(doc, "M.2220-1，§4.1，PDF第32–33页。")

    add_heading(doc, "12.2 三个TACAN台站对高精度CDMA接收机", 3)
    add_formula(
        doc,
        "PDCLIM,j = 2 × (3.5 μs + 1 μs) × 3600 s⁻¹ = 0.0324",
        "每个台站发射脉冲对，因此乘以2。",
    )
    add_formula(
        doc,
        "PDCLIM,agg = 1 − (1 − 0.0324)³ = 0.09408",
        "三个同类强脉冲源的净时间并集。",
    )

    add_heading(doc, "12.3 单部监视雷达对高精度半无代码接收机", 3)
    add_formula(
        doc,
        "PDCLIM = (50 μs + 1 μs) × 1500 s⁻¹ = 0.0765",
        "示例假设脉冲峰值高于饱和门限但低于生存门限。",
    )

    add_heading(doc, "12.4 星载SAR对E6-BC/L6接收机", 3)
    add_table(
        doc,
        ["参数", "报告示例值"],
        [
            ("SAR主瓣EIRP", "74.5 dBW"),
            ("轨道高度", "628 km"),
            ("中心频率", "1257.5 MHz"),
            ("带宽", "28 MHz"),
            ("脉宽/PRF组合", "11 μs / 6000 s⁻¹"),
            ("接收机恢复时间", "1 μs"),
            ("PDCLIM", "(11+1)×6000×10⁻⁶ = 0.072"),
        ],
        widths_cm=[6.0, 9.6],
    )
    add_source_note(doc, "M.2220-1，§4.2–4.3，PDF第34–37页。")

    add_heading(doc, "13. M.2220-1得出了什么结论", 2)
    add_bullet(doc, "可以用统一方法计算带内和邻带脉冲源对航空和地面RNSS接收机的聚合参数。")
    add_bullet(doc, "PDC和RI必须同时考虑：强脉冲造成时间损失，弱脉冲造成等效噪声。")
    add_bullet(doc, "多源PDC必须使用碰撞/补集模型，不能简单相加。")
    add_bullet(doc, "接收机恢复时间、滤波器选择性和熄减/饱和结构会显著改变结果。")
    add_bullet(doc, "计算得到的聚合参数应送入M.2030，M.2220本身不提供最终兼容性判决。")
    add_bullet(doc, "对更高RNSS频段的适用性仍需进一步研究。")

    add_heading(doc, "14. M.2220-1的局限及项目使用注意事项", 2)
    add_table(
        doc,
        ["局限", "对项目的影响", "建议处理"],
        [
            ("部分发射机数据库和运行密度具有地区性", "不能直接代表项目所在区域", "建立本地区台站、功率、频率和工作率数据库"),
            ("接收机模型是代表性模型", "真实设备门限和恢复时间可能差异很大", "进行接收机射频注入试验"),
            ("矩形/等效脉宽近似", "复杂chirp和滤波器裙边可能被简化", "使用实测发射谱与滤波响应积分"),
            ("静态或统计场景不等于轨道同时照射", "无法给出多星事件时间窗口", "与M.2305轨道级仿真结合"),
            ("输出为中间参数", "不能直接报告“合格/不合格”", "继续使用M.2030和任务性能指标"),
        ],
        widths_cm=[4.3, 5.2, 6.1],
        font_size=8.4,
    )


def add_m2305_section(doc: Document, charts: dict[str, Path]) -> None:
    add_heading(doc, "第二篇：Report ITU-R M.2305-1详细解读", 1)

    add_heading(doc, "15. 报告要解决什么问题", 2)
    add_body(
        doc,
        "M.2305-1关注的不是单颗SAR的瞬时最坏峰值，而是多个EESS（有源）传感器在真实轨道和波束几何下，是否会在同一时刻共同照射同一个地面RNSS接收机。"
        "报告统计的是同时照射事件的次数和持续时间，而不是直接给出最终C/N0劣化。",
    )
    add_callout(
        doc,
        "报告边界",
        "M.2305-1回答“事件会不会发生、发生多频繁、持续多久”。"
        "它不独立回答“事件是否可以接受”。后者需要将实际EESS参数送入M.2220和M.2030。",
        GREEN_LIGHT,
    )
    add_source_note(doc, "M.2305-1，§1–§2及§6，PDF第3–5页、第18–19页。")

    add_heading(doc, "16. 报告如何定义聚合事件", 2)
    add_body(
        doc,
        "报告把每颗传感器在接收机端的峰值功率与一个特定强脉冲门限比较。主分析使用−129 dBW门限。"
        "在某一时刻，如果两颗EESS传感器都高于该门限，就记录为强—强同时照射事件。",
    )
    add_table(
        doc,
        ["组合", "含义", "后续需要计算"],
        [
            ("强—强", "两颗传感器均超过饱和/熄减门限", "聚合PDC；M.2305的主要统计对象"),
            ("强—弱", "一颗超过门限，另一颗低于门限", "强源PDC + 弱源RI；弱源可能被强脉冲时间遮蔽"),
            ("弱—强", "与上项对称", "PDC与RI联合计算"),
            ("弱—弱", "两颗都低于门限", "RI功率密度求和"),
        ],
        widths_cm=[2.8, 6.2, 6.6],
    )
    add_body(
        doc,
        "因此，“没有强—强事件”并不等于完全没有聚合干扰；弱—弱和强—弱仍可能提高有效噪声。"
        "报告为了聚焦轨道重叠，把主要统计限制在强—强事件。",
    )

    add_heading(doc, "17. 三类EESS传感器模型", 2)
    add_table(
        doc,
        ["参数", "SAR3", "SAR6", "Scatterometer 2"],
        [
            ("轨道", "757 km，98°，18:00太阳同步", "628 km，97.9°，12:00太阳同步", "680 km，98°，18:00太阳同步"),
            ("峰值EIRP", "68.4 dBW", "74.5 dBW", "约60 dBW"),
            ("天线波束", "约25°×0.8°，窄方位", "4.6°×1.3°，宽波束SAR示例", "2.6°×2.6°旋转点波束"),
            ("观察角", "30°", "研究中取59°", "34°，方位360°旋转"),
            ("频率/带宽", "1215–1300 MHz，78 MHz", "1236.5/1257.5/1278.5 MHz，28 MHz", "1215–1300 MHz，2×1 MHz"),
            ("脉宽/PRF", "78 μs / 2400 Hz", "18–43 μs / 1550–3640 Hz", "15 μs / 3500 Hz"),
            ("发射占空比", "18.7%", "6.81%", "10.5%"),
            ("扫描", "固定侧视", "固定侧视", "13–14.6 rpm，约4.1–4.6 s一圈"),
        ],
        widths_cm=[3.0, 4.3, 4.3, 4.3],
        font_size=7.9,
    )
    add_body(
        doc,
        "三个模型故意代表不同的空间覆盖特性：SAR3具有很窄的方位波束；SAR6的示例覆盖更宽；散射计波束很窄但快速旋转，会在一次过境中多次扫过接收机。"
        "这使事件统计既受轨道交叉影响，也受波束驻留方式影响。",
    )

    add_heading(doc, "18. 仿真系统怎样搭建", 2)
    add_numbered(doc, "在STK中建立太阳同步轨道卫星对象，并赋予轨道高度、倾角和地方时升交点。")
    add_numbered(doc, "为每颗卫星添加发射机、射频参数、观察角和天线方向图。")
    add_numbered(doc, "在地面设置RNSS接收机点，纬度为40°、50°、60°、70°和80°，主要经度为75°W。")
    add_numbered(doc, "计算每颗传感器单独作用时的Received Isotropic Power（RIP）时间序列。")
    add_numbered(doc, "在MATLAB中加入RNSS接收天线仰角增益，把RIP转换为接收机端干扰功率。")
    add_numbered(doc, "逐时刻统计超过−129 dBW门限的传感器数量，并提取两颗同时超限的事件起止时间。")
    add_numbered(doc, "对120天结果统计事件次数、最大持续时间和中位持续时间。")
    add_source_note(doc, "M.2305-1，§4，PDF第9–15页。")

    add_heading(doc, "19. 为什么采用不同时间步长", 2)
    add_table(
        doc,
        ["场景", "时间步长", "理由"],
        [
            ("SAR—SAR", "10 s", "SAR照射足迹在地面移动较慢，单次照射持续数分钟"),
            ("SAR—散射计", "先找共同过境窗口，再用50 ms对每个10 min窗口精细仿真", "散射计约4.1 s旋转一周，扫过接收机的事件很短"),
            ("步长验证", "用10 ms结果抽查50 ms", "确认50 ms能够捕获强脉冲门限超限事件；峰值可能存在约数dB差异"),
        ],
        widths_cm=[3.7, 5.7, 6.2],
    )
    add_body(
        doc,
        "这体现了动态干扰仿真的基本原则：时间步长必须由最快的几何或波束变化决定，而不是统一采用一个方便的数值。"
        "步长过大会漏掉旋转波束事件；步长过小则使120天仿真计算量不可接受。",
    )

    add_heading(doc, "20. Study A：规划/既有传感器组合", 2)
    add_table(
        doc,
        ["场景", "组合", "轨道关系", "研究目的"],
        [
            ("A1", "SAR3 + SAR6", "正交太阳同步轨道", "窄波束与宽波束SAR重叠"),
            ("A2", "SAR3 + Scatterometer 2", "近共面，不同高度", "SAR固定波束与旋转散射计多次扫描"),
            ("A3", "SAR6 + Scatterometer 2", "正交轨道", "宽波束SAR与旋转点波束重叠"),
        ],
        widths_cm=[2.2, 5.0, 4.2, 4.9],
    )
    add_table(
        doc,
        ["纬度", "A1事件", "A2扫描事件（过境）", "A3扫描事件（过境）"],
        [
            ("40°", "0", "91（5）", "0"),
            ("50°", "0", "82（4）", "0"),
            ("60°", "0", "87（5）", "0"),
            ("70°", "0", "139（9）", "6（1）"),
            ("80°", "41", "409（18）", "1186（63）"),
        ],
        widths_cm=[2.5, 3.8, 5.0, 5.0],
    )
    add_figure(
        doc,
        charts["study_a"],
        "图2  M.2305-1 Study A在120天内的强—强事件计数（纵轴采用对称对数尺度；散射计列按单次扫描计数）",
    )
    add_body(
        doc,
        "A1只在80°高纬出现41次重叠，最长6.5 min、中位1.5 min。A2和A3由于散射计在一次共同过境中反复扫描，表中的大数是“波束扫描事件”而非轨道过境次数；括号内才是共同过境次数。"
        "例如A3在80°有1186次扫描超限，但对应63次轨道共同过境。",
    )
    add_callout(
        doc,
        "Study A不能直接解读为运行风险",
        "报告没有把传感器实际开机占空比和“同一区域不得同时成像”的运营约束全部计入。"
        "此外，所列任务寿命不一定重叠。因此这些数字用于识别几何潜势，不等于实际年度超限次数。",
        ORANGE_LIGHT,
    )

    add_heading(doc, "21. Study B：未来部署的敏感性研究", 2)
    add_table(
        doc,
        ["场景", "模型", "关键问题"],
        [
            ("B4", "两颗宽波束SAR，正交轨道", "高纬轨道交叉会形成多少事件"),
            ("B5", "两颗宽波束SAR，共面轨道", "最不利轨道布置是否使所有纬度均出现事件"),
            ("B6", "两颗EESS仅做视距可见性研究", "不考虑天线抑制时的几何上限"),
        ],
        widths_cm=[2.2, 6.0, 7.9],
    )
    add_table(
        doc,
        ["纬度", "B4事件", "B5事件", "B6视距事件"],
        [
            ("40°", "0", "44", "0"),
            ("50°", "0", "55", "25"),
            ("60°", "0", "68", "185"),
            ("70°", "32", "116", "296"),
            ("80°", "194", "189", "459"),
        ],
        widths_cm=[3.0, 4.3, 4.3, 4.3],
    )
    add_table(
        doc,
        ["纬度", "B4最大/中位", "B5最大/中位", "B6最大/中位"],
        [
            ("40°", "—", "10.2/2.1 min", "—"),
            ("50°", "—", "10.2/2.2 min", "6.1/3.6 min"),
            ("60°", "—", "10.0/2.3 min", "11.9/5.2 min"),
            ("70°", "7.0/1.6 min", "9.8/2.3 min", "13.3/5.4 min"),
            ("80°", "8.2/2.3 min", "10.3/3.3 min", "13.4/6.6 min"),
        ],
        widths_cm=[3.0, 4.3, 4.3, 4.3],
    )
    add_figure(
        doc,
        charts["study_b"],
        "图3  M.2305-1 Study B在120天内的事件计数（B6为视距几何上限，并非所有事件都超过−129 dBW）",
    )
    add_body(
        doc,
        "B4表明，正交宽波束SAR的强—强事件主要集中在70°以上。B5表明，如果两颗宽波束传感器近共面，所有研究纬度都可能发生事件，且单次最长约10 min。"
        "因此，在考虑天线抑制的场景中，B5是报告认定的最不利配置。B6说明纯视距重叠从50°起就可能很频繁，但视距存在并不等价于接收功率一定超过强脉冲门限。",
    )

    add_heading(doc, "22. 附件2：单颗分别合格，聚合后仍可能超限", 2)
    add_table(
        doc,
        ["系统", "单独PDC/结果", "单独劣化"],
        [
            ("SAR1", "PDC≈0.00225", "0.019 dB"),
            ("假设SARA", "PDC≈0.0223", "0.196 dB"),
            ("同时作用", "按M.2030聚合", "0.215 dB"),
        ],
        widths_cm=[4.5, 5.5, 5.5],
    )
    add_figure(
        doc,
        charts["annex2"],
        "图4  M.2305-1附件2示例：两个系统单独均低于0.2 dB，但同时作用达到0.215 dB",
        width_cm=14.3,
    )
    add_body(
        doc,
        "这个例子不是说所有SAR劣化可以永远简单相加，而是针对报告所采用的硬限幅/强脉冲模型，证明新增系统必须在“当前累计脉冲环境”上评估。"
        "因此，项目评审不能逐颗卫星分别签字后就认定星座整体合格。",
    )
    add_source_note(doc, "M.2305-1，附件2，PDF第20–23页。")

    add_heading(doc, "23. M.2305-1得出了什么结论", 2)
    add_bullet(doc, "多个EESS有源传感器同时照射同一RNSS接收机的事件确实可能发生，其次数和持续时间强烈依赖纬度、轨道面关系和波束宽度。")
    add_bullet(doc, "正交太阳同步轨道具有较大经度分离，事件更集中于高纬；非正交或共面轨道可能更不利。")
    add_bullet(doc, "散射计旋转波束会在一次共同过境中形成多次短事件，因此必须区分“扫描事件数”和“轨道过境数”。")
    add_bullet(doc, "真实评估必须使用实际三维天线方向图、实际轨道、运行占空比和照射区域。")
    add_bullet(doc, "M.2305只识别事件潜势；最终量化应使用M.2220和M.2030。")
    add_bullet(doc, "即使各单颗SAR独立评估均合格，同时作用仍可能使累积劣化超限。")

    add_heading(doc, "24. M.2305-1的局限及项目使用注意事项", 2)
    add_table(
        doc,
        ["局限", "可能造成的偏差", "工程补充"],
        [
            ("只重点研究两颗传感器", "无法直接给出三颗及以上的联合概率", "用全部规划/在轨星历开展多源Monte Carlo或事件扫描"),
            ("仅研究18:00和12:00等轨道关系", "非正交轨道可能更不利", "覆盖所有轨道面和地方时组合"),
            ("天线由两个主平面切面合成", "非主平面旁瓣可能不准确", "导入实测或全波三维方向图"),
            ("Study A未计实际开机率和海洋禁发", "事件数可能偏保守", "加入任务计划、成像区域和模式调度"),
            ("−129 dBW是特定门限", "不适用于所有RNSS接收机", "用实测饱和/熄减门限替换"),
            ("主要统计强—强事件", "弱—弱、强—弱贡献可能遗漏", "逐时刻同时计算PDC与RI"),
        ],
        widths_cm=[4.2, 5.2, 6.2],
        font_size=8.3,
    )


def add_integrated_workflow(doc: Document) -> None:
    add_heading(doc, "第三部分：两份报告怎样组合用于详细分析", 1)

    add_heading(doc, "25. 一套可执行的端到端流程", 2)
    steps = [
        "建立所有脉冲源清单：地面雷达、机载源、星载SAR/散射计及当前连续干扰基线。",
        "为每个源建立发射功率、实测频谱、脉宽、PRF、天线三维方向图、位置/轨道和工作时序。",
        "建立RNSS接收机模型：天线、RF/IF滤波器、饱和/熄减门限、恢复时间、NLIM、AGC/ADC和跟踪模式。",
        "按M.2305思路做轨道—姿态—波束动态仿真，得到每个源的接收功率时间序列。",
        "逐时刻把接收功率与门限比较，按M.2220将每个源分类为强脉冲或弱脉冲。",
        "根据真实频谱重叠计算PW_eff；对强源计算PDC，对弱源计算RI。",
        "按补集乘积和功率密度求和形成PDC_agg与RI_agg。",
        "把逐时刻聚合参数输入M.2030，计算新增劣化D(t)。",
        "统计最大劣化、超限总时长、单次最长超限、事件次数、纬度/区域分布和多星同时照射次数。",
        "对关键工况开展射频注入试验，校准接收机门限、恢复时间、真实滤波和瞬时C/N0响应。",
        "设计中心频率、脉宽、PRF、波束、区域和运行时序缓解方案，重新仿真并形成闭环。",
    ]
    for step in steps:
        add_numbered(doc, step)

    add_heading(doc, "26. 项目最小参数清单", 2)
    add_table(
        doc,
        ["类别", "至少需要的参数"],
        [
            ("脉冲源", "中心频率、峰值/平均功率、脉宽、PRF、占空比、调频带宽、发射谱、极化、工作率"),
            ("天线与几何", "完整三维增益、扫描规律、观察角、轨道/位置、姿态、星地斜距、到达角"),
            ("传播", "自由空间损耗、极化失配、地形/遮挡、必要时异常传播"),
            ("RNSS接收机", "天线方向图、RF/IF滤波、系统噪声、熄减/饱和门限、生存电平、τr、NLIM、AGC/ADC"),
            ("运行环境", "已有脉冲源基线PDC/RI、连续干扰I0,WB/N0、任务区域、开机时段、其他卫星计划"),
            ("验收输出", "峰值功率、PDC/RI时序、D(t)、CDF/CCDF、超限持续时间、事件地图、试验测得C/N0/失锁/恢复"),
        ],
        widths_cm=[3.5, 12.2],
    )

    add_heading(doc, "27. 两份报告最值得项目借鉴的方法", 2)
    add_table(
        doc,
        ["方法", "来自哪份报告", "项目落地方式"],
        [
            ("强/弱脉冲分流", "M.2220", "按实测接收机门限逐时刻分类，避免全部当噪声或全部当饱和"),
            ("净PDC聚合", "M.2220", "使用时间并集而非算术和，处理多源脉冲碰撞"),
            ("chirp有效脉宽", "M.2220", "用发射谱与接收滤波器的重叠积分替代名义带宽判断"),
            ("源—路径—接收机模块化", "M.2220", "软件中分离发射机数据库、传播、接收机模型，便于复用和审查"),
            ("长时轨道事件统计", "M.2305", "至少覆盖轨道重复周期或代表性长期窗口"),
            ("多尺度时间步长", "M.2305", "慢SAR用秒级，旋转波束用毫秒级窗口，兼顾精度与计算量"),
            ("事件次数与持续时间分开", "M.2305", "避免只报告最大功率，补充单次持续、重访间隔和累计时间"),
            ("单源合格不代表聚合合格", "M.2305附件2", "所有新增系统必须在现有干扰基线上评估"),
        ],
        widths_cm=[3.8, 3.2, 8.7],
        font_size=8.3,
    )

    add_heading(doc, "28. 常见误解与正确表述", 2)
    add_table(
        doc,
        ["误解", "正确理解"],
        [
            ("PDC就是发射机占空比", "PDC是接收机端有效时间损失，包含滤波后的有效脉宽和恢复时间。"),
            ("只要两颗卫星可见就一定超限", "可见性只是几何上限，还要计算天线增益、频谱重叠和接收机门限。"),
            ("没有强—强事件就没有聚合影响", "弱—弱和强—弱仍会通过RI和PDC共同产生劣化。"),
            ("单颗分别低于0.2 dB即可批准", "必须加入现有和规划脉冲源的累计影响。"),
            ("报告中的−129 dBW、1 μs可直接用于所有设备", "这些是特定模型参数，详细分析应由实测设备参数替换。"),
            ("M.2305给出了最终兼容性结论", "它给事件潜势；最终接受性仍需M.2220/M.2030及接收机试验。"),
        ],
        widths_cm=[6.0, 9.7],
    )

    add_heading(doc, "29. 总结", 2)
    add_body(
        doc,
        "M.2220-1的贡献是把复杂、异质的脉冲环境变成结构统一、可进入性能公式的聚合参数。它解决了“多个脉冲源怎样算在一起”的问题。"
        "M.2305-1的贡献是把单源静态最坏情况扩展到轨道、纬度、波束和时间维度，解决了“多个星载源什么时候真的会同时出现”的问题。",
    )
    add_body(
        doc,
        "两份报告结合后，详细分析应当同时保留功率域、频率域、时间域和几何域信息。"
        "最终工程结论不应只有一个最大I/N或一个0.2 dB比较值，而应包括事件地图、持续时间、累计概率、最坏波形、接收机实测参数以及缓解后的复算结果。",
    )

    add_heading(doc, "附录A：核心公式速查", 1)
    add_table(
        doc,
        ["公式", "用途"],
        [
            ("PDC_j=(PW_j,eff+τr)PRF_j", "单个强脉冲源的有效失效时间比例"),
            ("PDC_agg=1−∏(1−PDC_j)", "多个强脉冲源的净聚合"),
            ("RI=Σ[P_i dc_i/(N0 BW)]", "多个弱脉冲源的等效噪声比"),
            ("PW_eff=PW(BW_overlap/BW_chirp)", "宽带chirp只部分进入接收通带时修正脉宽"),
            ("PRF_eff=PRF(τobs/Tc)", "特定极窄环路接收机的快速扫描折减；不能普遍采用"),
            ("D=10log10(N0,EFF+Y/N0,EFF)", "把新增源转换为相对基线的性能劣化"),
        ],
        widths_cm=[7.1, 8.6],
    )

    add_heading(doc, "附录B：官方原文入口", 1)
    p = doc.add_paragraph()
    r = p.add_run("Report ITU-R M.2220-1：")
    set_run_font(r, bold=True)
    add_hyperlink(p, "官方英文PDF", URL_M2220)
    set_paragraph_spacing(p, after=6)

    p = doc.add_paragraph()
    r = p.add_run("Report ITU-R M.2305-1：")
    set_run_font(r, bold=True)
    add_hyperlink(p, "官方英文PDF", URL_M2305)
    set_paragraph_spacing(p, after=6)

    p = doc.add_paragraph()
    r = p.add_run("Recommendation ITU-R M.2030：")
    set_run_font(r, bold=True)
    add_hyperlink(p, "ITU官方版本页面", URL_M2030)
    set_paragraph_spacing(p, after=6)

    p = doc.add_paragraph()
    r = p.add_run("Report ITU-R RS.2537-0：")
    set_run_font(r, bold=True)
    add_hyperlink(p, "ITU官方版本页面", URL_RS2537)
    set_paragraph_spacing(p, after=6)

    add_callout(
        doc,
        "证据边界",
        "正文中的报告结构、参数、算例和结论均来自两份ITU-R原文。"
        "“项目落地流程”“最小参数清单”和部分解释性表述属于基于原文方法的工程化整理，不应被表述为ITU-R逐字规定。",
        LIGHT_GRAY,
    )


def build_document() -> Path:
    charts = generate_charts()
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_front_matter(doc)
    add_m2220_section(doc, charts)
    doc.add_page_break()
    add_m2305_section(doc, charts)
    doc.add_page_break()
    add_integrated_workflow(doc)

    # Ask Word to update fields when opened.
    settings = doc.settings._element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)

    doc.save(DOCX_PATH)
    return DOCX_PATH


if __name__ == "__main__":
    output = build_document()
    print(output)
