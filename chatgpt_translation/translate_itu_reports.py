#!/usr/bin/env python3
from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import re
import statistics
import sys
import time
from pathlib import Path
from typing import Any, Iterable

import fitz  # PyMuPDF
import requests
import torch
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from transformers import AutoModelForSeq2SeqLM, AutoTokenizer

REPORTS: dict[str, dict[str, str]] = {
    "SA.2425": {
        "number": "Report ITU-R SA.2425-0",
        "date": "09/2018",
        "url": "https://www.itu.int/dms_pub/itu-r/opb/rep/R-REP-SA.2425-2018-PDF-E.pdf",
        "filename": "R-REP-SA.2425-2018-PDF-E.pdf",
        "zh_title": "满足短期任务非对地静止卫星空间操作业务频谱需求的研究",
        "en_title": "Studies to accommodate spectrum requirements in the space operation service for non-geostationary satellites with short duration missions",
        "output": "ITU-R_SA.2425-0_中文翻译.docx",
    },
    "SA.2426": {
        "number": "Report ITU-R SA.2426-0",
        "date": "09/2018",
        "url": "https://www.itu.int/dms_pub/itu-r/opb/rep/R-REP-SA.2426-2018-PDF-E.pdf",
        "filename": "R-REP-SA.2426-2018-PDF-E.pdf",
        "zh_title": "1 GHz以下短期任务非对地静止卫星空间操作业务遥测、跟踪和遥控的技术特性",
        "en_title": "Technical characteristics for telemetry, tracking and command in the space operation service below 1 GHz for non-GSO satellites with short duration missions",
        "output": "ITU-R_SA.2426-0_中文翻译.docx",
    },
    "SA.2427": {
        "number": "Report ITU-R SA.2427-0",
        "date": "09/2018",
        "url": "https://www.itu.int/dms_pub/itu-r/opb/rep/R-REP-SA.2427-2018-PDF-E.pdf",
        "filename": "R-REP-SA.2427-2018-PDF-E.pdf",
        "zh_title": "1 GHz以下空间操作业务现有频率划分适用性研究及可能新增和/或升级频率划分的补充共用研究",
        "en_title": "Studies on the suitability of existing allocations to the space operation service below 1 GHz and additional sharing studies on possible new and/or upgraded allocations",
        "output": "ITU-R_SA.2427-0_中文翻译.docx",
    },
}

# Longest phrases first. The placeholders are restored after machine translation so that
# ITU terminology remains consistent throughout the documents.
GLOSSARY: list[tuple[str, str]] = [
    ("non-geostationary satellites with short duration missions", "短期任务非对地静止卫星"),
    ("non-GSO satellites with short duration missions", "短期任务非GSO卫星"),
    ("NGSO satellites with short duration missions", "短期任务非GSO卫星"),
    ("non-geostationary-satellite orbit", "非对地静止卫星轨道"),
    ("non-geostationary satellite orbit", "非对地静止卫星轨道"),
    ("World Radiocommunication Conference", "世界无线电通信大会"),
    ("Radiocommunication Assembly", "无线电通信全会"),
    ("Radiocommunication Sector", "无线电通信部门"),
    ("International Telecommunication Union", "国际电信联盟"),
    ("telemetry, tracking and command", "遥测、跟踪和遥控"),
    ("space operation service", "空间操作业务"),
    ("Earth exploration-satellite service", "卫星地球探测业务"),
    ("meteorological-satellite service", "气象卫星业务"),
    ("meteorological aids service", "气象辅助业务"),
    ("mobile-satellite service", "移动卫星业务"),
    ("amateur-satellite service", "业余卫星业务"),
    ("aeronautical mobile service", "航空移动业务"),
    ("maritime mobile service", "水上移动业务"),
    ("radio astronomy service", "射电天文业务"),
    ("space research service", "空间研究业务"),
    ("radiolocation service", "无线电定位业务"),
    ("radionavigation service", "无线电导航业务"),
    ("fixed service", "固定业务"),
    ("mobile service", "移动业务"),
    ("space-to-Earth", "空对地"),
    ("Earth-to-space", "地对空"),
    ("power flux-density", "功率通量密度"),
    ("power flux density", "功率通量密度"),
    ("power spectral density", "功率谱密度"),
    ("equivalent isotropically radiated power", "等效全向辐射功率"),
    ("necessary bandwidth", "必要带宽"),
    ("occupied bandwidth", "占用带宽"),
    ("out-of-band emissions", "带外发射"),
    ("out-of-band emission", "带外发射"),
    ("spurious emissions", "杂散发射"),
    ("spurious emission", "杂散发射"),
    ("unwanted emissions", "非必要发射"),
    ("unwanted emission", "非必要发射"),
    ("sharing and compatibility studies", "共用与兼容性研究"),
    ("sharing studies", "共用研究"),
    ("compatibility studies", "兼容性研究"),
    ("protection criteria", "保护准则"),
    ("protection criterion", "保护准则"),
    ("short duration mission", "短期任务"),
    ("short-duration mission", "短期任务"),
    ("earth station", "地球站"),
    ("space station", "空间电台"),
    ("link budget", "链路预算"),
    ("link margin", "链路余量"),
    ("system noise temperature", "系统噪声温度"),
    ("noise temperature", "噪声温度"),
    ("antenna radiation pattern", "天线辐射方向图"),
    ("antenna gain", "天线增益"),
    ("minimum elevation angle", "最小仰角"),
    ("duty cycle", "占空比"),
    ("primary allocation", "主要业务划分"),
    ("secondary allocation", "次要业务划分"),
    ("frequency allocation", "频率划分"),
    ("frequency band", "频段"),
]

ZH_REPLACEMENTS = {
    "空间运营业务": "空间操作业务",
    "空间运行服务": "空间操作业务",
    "空间运营服务": "空间操作业务",
    "遥测、跟踪和命令": "遥测、跟踪和遥控",
    "遥测、跟踪和指令": "遥测、跟踪和遥控",
    "非地球静止": "非对地静止",
    "非地球同步": "非对地静止",
    "地球到空间": "地对空",
    "空间到地球": "空对地",
    "功率光谱密度": "功率谱密度",
    "等效各向同性辐射功率": "等效全向辐射功率",
    "外带发射": "带外发射",
    "虚假发射": "杂散发射",
    "保护标准": "保护准则",
}

CJK_FONT = "Noto Sans CJK SC"
LATIN_FONT = "Arial"


def log(message: str) -> None:
    print(time.strftime("[%H:%M:%S]"), message, flush=True)


def download(url: str, destination: Path) -> None:
    if destination.exists() and destination.stat().st_size > 50_000:
        log(f"Using cached PDF: {destination}")
        return
    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/140 Safari/537.36",
        "Accept": "application/pdf,*/*;q=0.8",
    }
    last: Exception | None = None
    for attempt in range(6):
        try:
            with requests.get(url, headers=headers, timeout=180, stream=True) as response:
                response.raise_for_status()
                with destination.open("wb") as handle:
                    for chunk in response.iter_content(chunk_size=1024 * 1024):
                        if chunk:
                            handle.write(chunk)
            head = destination.read_bytes()[:5]
            if head != b"%PDF-":
                raise RuntimeError(f"Downloaded file is not a PDF: {head!r}")
            log(f"Downloaded {destination.name}: {destination.stat().st_size:,} bytes")
            return
        except Exception as exc:  # noqa: BLE001
            last = exc
            time.sleep(3 * (attempt + 1))
    raise RuntimeError(f"Unable to download {url}: {last!r}")


def normalize_source(text: str) -> str:
    text = text.replace("\u00ad", "")
    text = text.replace("\ufb01", "fi").replace("\ufb02", "fl")
    text = text.replace("\u2011", "-").replace("\u2010", "-")
    text = text.replace("\u2212", "-")
    text = re.sub(r"(?<=\w)-\s*\n\s*(?=\w)", "", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def is_header_or_footer(text: str, page_index: int, y0: float, y1: float, page_height: float) -> bool:
    stripped = " ".join(text.split())
    if not stripped:
        return True
    if page_index > 0 and y0 < 65:
        if re.match(r"^(?:\d+\s+)?(?:Rep\.|Report|REPORT)\s+ITU-R\s+SA\.242[567]", stripped, re.I):
            return True
        if re.match(r"^Rep\.\s+ITU-R\s+SA\.242[567]-0\s+\d+$", stripped, re.I):
            return True
    if y1 > page_height - 42 and re.fullmatch(r"\d+", stripped):
        return True
    return False


def table_like(lines: list[str]) -> bool:
    if len(lines) < 4:
        return False
    short = sum(1 for line in lines if len(line) <= 35)
    numerical = sum(1 for line in lines if re.search(r"\d", line))
    separators = sum(1 for line in lines if "  " in line or "\t" in line)
    return (short / len(lines) > 0.55 and numerical / len(lines) > 0.25) or separators > 1


def heading_level(text: str, font_size: float, median_size: float) -> int | None:
    compact = " ".join(text.split())
    if re.match(r"^(ANNEX|APPENDIX)\b", compact, re.I):
        return 1
    if re.match(r"^(REPORT|RECOMMENDATION)\s+ITU-R\b", compact, re.I):
        return 1
    match = re.match(r"^(\d+(?:\.\d+){0,4})\s+\D", compact)
    if match:
        return min(3, match.group(1).count(".") + 1)
    if font_size >= median_size + 3.0:
        return 1
    if font_size >= median_size + 1.4 and len(compact) < 180:
        return 2
    if len(compact) < 150 and compact.isupper() and re.search(r"[A-Z]{3}", compact):
        return 2
    return None


def logical_units(block_text: str, kind_hint: str) -> list[str]:
    lines = [re.sub(r"\s+", " ", line).strip() for line in block_text.splitlines()]
    lines = [line for line in lines if line]
    if not lines:
        return []
    if kind_hint == "table" or table_like(lines):
        return lines
    units: list[str] = []
    current = ""
    for line in lines:
        force_new = bool(
            re.match(r"^(?:\d+(?:\.\d+){0,4}\s+|[-•–]\s+|TABLE\s+\S+|FIGURE\s+\S+|NOTE\s*[–-])", line, re.I)
        )
        if force_new and current:
            units.append(current.strip())
            current = ""
        if not current:
            current = line
        else:
            if current.endswith((".", ":", ";", "?", "!", ")")) and len(current) > 80:
                units.append(current.strip())
                current = line
            else:
                current += " " + line
    if current:
        units.append(current.strip())
    return units


def extract_units(pdf_path: Path) -> tuple[list[dict[str, Any]], int]:
    document = fitz.open(pdf_path)
    all_units: list[dict[str, Any]] = []
    for page_index, page in enumerate(document):
        raw = page.get_text("dict", sort=True)
        sizes: list[float] = []
        for block in raw.get("blocks", []):
            if block.get("type") != 0:
                continue
            for line in block.get("lines", []):
                for span in line.get("spans", []):
                    if span.get("text", "").strip():
                        sizes.append(float(span.get("size", 10.0)))
        median_size = statistics.median(sizes) if sizes else 10.0
        page_units: list[dict[str, Any]] = []
        for block in raw.get("blocks", []):
            if block.get("type") != 0:
                continue
            line_texts: list[str] = []
            font_sizes: list[float] = []
            bold = False
            for line in block.get("lines", []):
                pieces: list[str] = []
                for span in line.get("spans", []):
                    value = str(span.get("text", ""))
                    if value:
                        pieces.append(value)
                        font_sizes.append(float(span.get("size", median_size)))
                        font_name = str(span.get("font", ""))
                        if "bold" in font_name.lower():
                            bold = True
                line_text = "".join(pieces).strip()
                if line_text:
                    line_texts.append(line_text)
            text = normalize_source("\n".join(line_texts))
            if not text:
                continue
            x0, y0, x1, y1 = block.get("bbox", (0.0, 0.0, 0.0, 0.0))
            if is_header_or_footer(text, page_index, float(y0), float(y1), float(page.rect.height)):
                continue
            max_size = max(font_sizes) if font_sizes else median_size
            level = heading_level(text, max_size, median_size)
            hint = "heading" if level else ("table" if table_like(text.splitlines()) else "body")
            for unit in logical_units(text, hint):
                compact = unit.strip()
                if not compact:
                    continue
                kind = "body"
                if level:
                    kind = f"heading{level}"
                elif re.match(r"^(TABLE|FIGURE)\s+", compact, re.I):
                    kind = "caption"
                elif hint == "table":
                    kind = "table"
                elif re.match(r"^[-•–]\s+", compact):
                    kind = "list"
                page_units.append({"page": page_index + 1, "text": compact, "kind": kind, "bold": bold})
        all_units.extend(page_units)
        if (page_index + 1) % 25 == 0 or page_index + 1 == len(document):
            log(f"Extracted page {page_index + 1}/{len(document)}; units={len(all_units)}")
    page_count = len(document)
    document.close()
    return all_units, page_count


def protect_glossary(text: str) -> tuple[str, dict[str, str]]:
    mapping: dict[str, str] = {}
    output = text
    counter = 0

    def add_placeholder(value: str) -> str:
        nonlocal counter
        token = f"ZXQPH{counter:04d}QXZ"
        counter += 1
        mapping[token] = value
        return token

    for source, target in sorted(GLOSSARY, key=lambda pair: len(pair[0]), reverse=True):
        pattern = re.compile(re.escape(source), re.I)
        while True:
            match = pattern.search(output)
            if not match:
                break
            token = add_placeholder(target)
            output = output[: match.start()] + token + output[match.end() :]

    patterns = [
        r"https?://\S+",
        r"\b(?:Recommendation|Report|Resolution)\s+ITU-R\s+[A-Z]+\.\d+(?:-\d+)?\b",
        r"\bITU-R\s+[A-Z]+\.\d+(?:-\d+)?\b",
        r"\bRR\s+No\.\s*\d+(?:\.\d+)*\b",
        r"\bWRC-\d{2}\b",
        r"\b[A-Z][A-Z0-9/()-]{1,14}\b",
    ]
    for expression in patterns:
        regex = re.compile(expression)
        while True:
            match = regex.search(output)
            if not match:
                break
            token = add_placeholder(match.group(0))
            output = output[: match.start()] + token + output[match.end() :]
    return output, mapping


def restore_placeholders(text: str, mapping: dict[str, str]) -> str:
    output = text
    for token, value in mapping.items():
        patterns = [
            re.escape(token),
            re.escape(" ".join(token)),
            re.escape(token.replace("PH", " PH ")),
        ]
        replaced = False
        for pattern in patterns:
            new_output, count = re.subn(pattern, value, output, flags=re.I)
            if count:
                output = new_output
                replaced = True
                break
        if not replaced:
            # The model occasionally inserts spaces inside placeholders. Match a permissive form.
            chars = r"\s*".join(re.escape(ch) for ch in token)
            output = re.sub(chars, value, output, flags=re.I)
    return output


def normalize_chinese(text: str) -> str:
    output = text.strip()
    for old, new in ZH_REPLACEMENTS.items():
        output = output.replace(old, new)
    output = output.replace(" ,", "，").replace(" ;", "；").replace(" :", "：")
    output = output.replace(" .", "。").replace(" ?", "？").replace(" !", "！")
    output = re.sub(r"(?<=[\u4e00-\u9fff])\s+(?=[\u4e00-\u9fff])", "", output)
    output = re.sub(r"\s+([，。；：！？、）】])", r"\1", output)
    output = re.sub(r"([（【])\s+", r"\1", output)
    output = re.sub(r"\s{2,}", " ", output)
    return output.strip()


def mostly_nonlinguistic(text: str) -> bool:
    alpha = len(re.findall(r"[A-Za-z]", text))
    if alpha < 3:
        return True
    nonspace = len(re.sub(r"\s", "", text))
    return nonspace > 0 and alpha / nonspace < 0.16 and bool(re.search(r"\d", text))


def split_long(text: str, tokenizer: Any, max_tokens: int = 380) -> list[str]:
    if len(tokenizer.encode(text, add_special_tokens=False)) <= max_tokens:
        return [text]
    clauses = re.split(r"(?<=[.!?;:])\s+", text)
    chunks: list[str] = []
    current = ""
    for clause in clauses:
        candidate = clause if not current else current + " " + clause
        if len(tokenizer.encode(candidate, add_special_tokens=False)) <= max_tokens:
            current = candidate
            continue
        if current:
            chunks.append(current)
            current = ""
        if len(tokenizer.encode(clause, add_special_tokens=False)) <= max_tokens:
            current = clause
        else:
            words = clause.split()
            piece = ""
            for word in words:
                candidate2 = word if not piece else piece + " " + word
                if len(tokenizer.encode(candidate2, add_special_tokens=False)) <= max_tokens:
                    piece = candidate2
                else:
                    if piece:
                        chunks.append(piece)
                    piece = word
            if piece:
                current = piece
    if current:
        chunks.append(current)
    return chunks or [text]


def translate_units(units: list[dict[str, Any]], model_dir: str, cache_path: Path) -> None:
    tokenizer = AutoTokenizer.from_pretrained(model_dir, local_files_only=False)
    model = AutoModelForSeq2SeqLM.from_pretrained(model_dir, local_files_only=False)
    model.eval()
    torch.set_num_threads(max(1, min(8, os.cpu_count() or 2)))

    cache: dict[str, str] = {}
    if cache_path.exists():
        try:
            cache = json.loads(cache_path.read_text(encoding="utf-8"))
            log(f"Loaded translation cache: {len(cache)} entries")
        except Exception:  # noqa: BLE001
            cache = {}

    jobs: list[tuple[int, str, str, dict[str, str]]] = []
    for index, unit in enumerate(units):
        source = unit["text"]
        key = hashlib.sha256(source.encode("utf-8")).hexdigest()
        if key in cache:
            unit["zh"] = cache[key]
            continue
        if mostly_nonlinguistic(source):
            unit["zh"] = source
            cache[key] = source
            continue
        protected, mapping = protect_glossary(source)
        parts = split_long(protected, tokenizer)
        unit["_translation_parts"] = len(parts)
        for part_index, part in enumerate(parts):
            jobs.append((index, f"{key}:{part_index}", part, mapping))

    log(f"Translation jobs: {len(jobs)} chunks; cached units: {sum('zh' in u for u in units)}")
    translated_parts: dict[int, list[str]] = {}
    batch_size = 24
    completed = 0
    for start in range(0, len(jobs), batch_size):
        batch = jobs[start : start + batch_size]
        sources = [item[2] for item in batch]
        encoded = tokenizer(
            sources,
            return_tensors="pt",
            padding=True,
            truncation=True,
            max_length=420,
        )
        with torch.inference_mode():
            generated = model.generate(
                **encoded,
                max_new_tokens=520,
                num_beams=1,
                do_sample=False,
                use_cache=True,
            )
        outputs = tokenizer.batch_decode(generated, skip_special_tokens=True)
        for item, output in zip(batch, outputs):
            unit_index, _, _, mapping = item
            restored = normalize_chinese(restore_placeholders(output, mapping))
            translated_parts.setdefault(unit_index, []).append(restored)
        completed += len(batch)
        if completed % 240 < batch_size or completed == len(jobs):
            log(f"Translated {completed}/{len(jobs)} chunks")

    for index, unit in enumerate(units):
        if "zh" in unit:
            continue
        pieces = translated_parts.get(index, [])
        translation = normalize_chinese(" ".join(piece for piece in pieces if piece))
        if not translation:
            translation = unit["text"]
        unit["zh"] = translation
        key = hashlib.sha256(unit["text"].encode("utf-8")).hexdigest()
        cache[key] = translation

    cache_path.write_text(json.dumps(cache, ensure_ascii=False, indent=2), encoding="utf-8")
    log(f"Saved translation cache: {cache_path}")


def set_cell_shading(cell: Any, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_run_font(run: Any, size: float, bold: bool = False, color: str | None = None) -> None:
    run.font.name = LATIN_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph: Any) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(end)
    set_run_font(run, 9, color="666666")


def configure_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = LATIN_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    normal.font.size = Pt(9.5)
    normal.paragraph_format.line_spacing = 1.18
    normal.paragraph_format.space_after = Pt(3)

    for name, size, color in [
        ("Title", 22, "1F4E79"),
        ("Subtitle", 11, "666666"),
        ("Heading 1", 15, "1F4E79"),
        ("Heading 2", 12.5, "2F5597"),
        ("Heading 3", 11, "365F91"),
    ]:
        style = styles[name]
        style.font.name = LATIN_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name != "Subtitle"
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(4)

    custom = {
        "SourcePage": (8.5, True, "FFFFFF"),
        "CaptionZH": (8.5, True, "44546A"),
        "TableTextZH": (8.2, False, "222222"),
        "TranslatorNote": (9, False, "666666"),
    }
    for style_name, (size, bold, color) in custom.items():
        if style_name not in styles:
            styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        style = styles[style_name]
        style.font.name = LATIN_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_after = Pt(2)
        style.paragraph_format.line_spacing = 1.05 if style_name == "TableTextZH" else 1.15
        if style_name == "TableTextZH":
            style.paragraph_format.left_indent = Cm(0.4)
            style.paragraph_format.right_indent = Cm(0.2)


def add_source_page_marker(document: Document, page_number: int) -> None:
    table = document.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, "4472C4")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.style = document.styles["SourcePage"]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(f"原文第 {page_number} 页")
    set_run_font(run, 8.5, bold=True, color="FFFFFF")


def make_docx(report: dict[str, str], units: list[dict[str, Any]], page_count: int, output_path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.7)
    configure_styles(document)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(report["zh_title"])
    set_run_font(run, 22, bold=True, color="1F4E79")

    number = document.add_paragraph()
    number.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = number.add_run(f"{report['number']}（{report['date']}）")
    set_run_font(run, 13, bold=True, color="2F5597")

    original = document.add_paragraph(style="Subtitle")
    original.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = original.add_run(report["en_title"])
    set_run_font(run, 10.5, color="666666")

    document.add_paragraph()
    note = document.add_paragraph(style="TranslatorNote")
    note.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = note.add_run(
        "译文说明：本文件为基于ITU-R英文原报告制作的中文参考译本，采用机器辅助翻译并统一了常用ITU术语。"
        "本译本不是国际电信联盟发布的官方中文文本；涉及法规效力、公式、数值、图表和最终工程结论时，应以英文原文为准。"
        "译文按照原报告页序组织，并标注对应原文页码，便于逐页核对。"
    )
    set_run_font(run, 9, color="666666")

    source = document.add_paragraph(style="TranslatorNote")
    source.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = source.add_run(f"原文来源：{report['url']}")
    set_run_font(run, 8.5, color="666666")

    document.add_page_break()

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header.add_run(f"{report['number']} - 中文参考译本")
    set_run_font(run, 8.5, color="666666")
    add_page_number(section.footer.paragraphs[0])

    current_page = None
    for unit in units:
        page_number = int(unit["page"])
        if page_number != current_page:
            current_page = page_number
            add_source_page_marker(document, page_number)
        text = str(unit.get("zh") or unit["text"]).strip()
        if not text:
            continue
        kind = unit.get("kind", "body")
        if kind.startswith("heading"):
            level = int(kind[-1])
            paragraph = document.add_paragraph(style=f"Heading {min(level, 3)}")
            paragraph.paragraph_format.keep_with_next = True
        elif kind == "caption":
            paragraph = document.add_paragraph(style="CaptionZH")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True
        elif kind == "table":
            paragraph = document.add_paragraph(style="TableTextZH")
        elif kind == "list":
            paragraph = document.add_paragraph(style="Normal")
            paragraph.paragraph_format.left_indent = Cm(0.5)
            paragraph.paragraph_format.first_line_indent = Cm(-0.25)
        else:
            paragraph = document.add_paragraph(style="Normal")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            paragraph.paragraph_format.first_line_indent = Cm(0.7)
        run = paragraph.add_run(text)
        size = 8.2 if kind == "table" else (8.5 if kind == "caption" else 9.5)
        set_run_font(run, size, bold=bool(unit.get("bold") and kind != "body"))

    # Add a compact end note with validation metadata.
    document.add_page_break()
    ending = document.add_paragraph(style="Heading 1")
    ending.add_run("译文校核信息")
    items = [
        f"英文原报告页数：{page_count} 页",
        f"译文内容单元数：{len(units)} 个",
        "编号、频率、功率、电平、百分比、公式符号及ITU-R文件编号原则上按原文保留。",
        "复杂表格在PDF文本抽取后可能按行展开；请结合标注的原文页码核对原表布局。",
    ]
    for item in items:
        paragraph = document.add_paragraph(style="Normal")
        paragraph.paragraph_format.left_indent = Cm(0.6)
        paragraph.paragraph_format.first_line_indent = Cm(-0.3)
        run = paragraph.add_run("• " + item)
        set_run_font(run, 9)

    core = document.core_properties
    core.title = f"{report['number']} 中文翻译"
    core.subject = report["zh_title"]
    core.author = "OpenAI"
    core.keywords = "ITU-R, SA, 中文翻译, 非GSO, 空间操作业务"
    core.comments = "非ITU官方译本；用于技术研读。"
    document.save(output_path)
    log(f"Saved DOCX: {output_path} ({output_path.stat().st_size:,} bytes)")


def verify_output(output_path: Path, units: list[dict[str, Any]], page_count: int, stats_path: Path) -> None:
    translated_chars = sum(len(str(unit.get("zh", ""))) for unit in units)
    cjk_chars = sum(len(re.findall(r"[\u4e00-\u9fff]", str(unit.get("zh", "")))) for unit in units)
    untranslated_alpha = sum(len(re.findall(r"[A-Za-z]", str(unit.get("zh", "")))) for unit in units)
    stats = {
        "output": str(output_path),
        "source_pages": page_count,
        "units": len(units),
        "translated_characters": translated_chars,
        "cjk_characters": cjk_chars,
        "latin_characters_remaining": untranslated_alpha,
        "docx_bytes": output_path.stat().st_size,
        "status": "PASS" if output_path.stat().st_size > 30_000 and cjk_chars > 1_000 else "FAIL",
    }
    stats_path.write_text(json.dumps(stats, ensure_ascii=False, indent=2), encoding="utf-8")
    log(json.dumps(stats, ensure_ascii=False))
    if stats["status"] != "PASS":
        raise RuntimeError(f"Output verification failed: {stats}")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--report", required=True, choices=sorted(REPORTS))
    parser.add_argument("--work-dir", default="work")
    parser.add_argument("--output-dir", default="out")
    parser.add_argument("--model", default="Helsinki-NLP/opus-mt-en-zh")
    args = parser.parse_args()

    report = REPORTS[args.report]
    work_dir = Path(args.work_dir).resolve() / args.report.replace(".", "_")
    output_dir = Path(args.output_dir).resolve()
    work_dir.mkdir(parents=True, exist_ok=True)
    output_dir.mkdir(parents=True, exist_ok=True)
    pdf_path = work_dir / report["filename"]
    download(report["url"], pdf_path)

    units_path = work_dir / "extracted_units.json"
    if units_path.exists():
        payload = json.loads(units_path.read_text(encoding="utf-8"))
        units = payload["units"]
        page_count = int(payload["page_count"])
        log(f"Loaded extracted units: {len(units)}")
    else:
        units, page_count = extract_units(pdf_path)
        units_path.write_text(
            json.dumps({"page_count": page_count, "units": units}, ensure_ascii=False), encoding="utf-8"
        )

    translate_units(units, args.model, work_dir / "translation_cache.json")
    translated_path = work_dir / "translated_units.json"
    translated_path.write_text(
        json.dumps({"page_count": page_count, "units": units}, ensure_ascii=False), encoding="utf-8"
    )

    output_path = output_dir / report["output"]
    make_docx(report, units, page_count, output_path)
    verify_output(output_path, units, page_count, output_dir / f"{args.report}_translation_stats.json")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
