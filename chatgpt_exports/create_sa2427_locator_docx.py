#!/usr/bin/env python3
from pathlib import Path
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path("generated/ITU-R_SA.2427-0原文页码定位与分析说明.docx")
OUT.parent.mkdir(parents=True, exist_ok=True)

BLUE = "1F4E79"
MID_BLUE = "4472C4"
LIGHT_BLUE = "D9EAF7"
VERY_LIGHT_BLUE = "EEF5FB"
GREY = "666666"
DARK = "222222"
WHITE = "FFFFFF"
ORANGE = "C55A11"


def set_run_font(run, size=10.5, bold=False, color=DARK):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="B4C7E7", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    separate = OxmlElement("w:fldChar"); separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])
    set_run_font(run, 9, color=GREY)


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, {1:15, 2:12.5, 3:11}[min(level,3)], True, {1:BLUE,2:MID_BLUE,3:BLUE}[min(level,3)])
    return p


def body(doc, text):
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0.74)
    r = p.add_run(text)
    set_run_font(r)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.first_line_indent = Cm(-0.35)
    r = p.add_run(text)
    set_run_font(r)
    return p


def numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.first_line_indent = Cm(-0.35)
    r = p.add_run(text)
    set_run_font(r)
    return p


def callout(doc, title, text, fill=VERY_LIGHT_BLUE, title_color=BLUE):
    t = doc.add_table(rows=1, cols=1)
    c = t.cell(0,0)
    set_cell_shading(c, fill)
    set_cell_border(c, MID_BLUE, "8")
    p = c.paragraphs[0]
    r = p.add_run(title)
    set_run_font(r, 10.5, True, title_color)
    p2 = c.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r2 = p2.add_run(text)
    set_run_font(r2, 10)
    return t


def table(doc, headers, rows, widths=None, header_fill=BLUE):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.autofit = False
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(c, header_fill)
        set_cell_border(c)
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(h)); set_run_font(r, 9.1, True, WHITE)
    for ridx, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cells[i])
            if ridx % 2:
                set_cell_shading(cells[i], "F8FBFE")
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(value)); set_run_font(r, 9)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


doc = Document()
sec = doc.sections[0]
sec.page_width = Cm(21.0); sec.page_height = Cm(29.7)
sec.top_margin = Cm(1.8); sec.bottom_margin = Cm(1.7)
sec.left_margin = Cm(2.0); sec.right_margin = Cm(2.0)
sec.header_distance = Cm(0.7); sec.footer_distance = Cm(0.7)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(10.5)
normal.paragraph_format.line_spacing = 1.35
normal.paragraph_format.space_after = Pt(5)
for name, size, color in [("Title",22,BLUE),("Subtitle",11.5,GREY),("Heading 1",15,BLUE),("Heading 2",12.5,MID_BLUE),("Heading 3",11,BLUE)]:
    st = styles[name]
    st.font.name = "Arial"
    st._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    st.font.size = Pt(size); st.font.bold = name != "Subtitle"
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.keep_with_next = True
    st.paragraph_format.space_before = Pt(9); st.paragraph_format.space_after = Pt(5)

# Cover
p = doc.add_paragraph(style="Title"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("ITU-R SA.2427-0\n原文页码定位与分析说明"); set_run_font(r,22,True,BLUE)
p = doc.add_paragraph(style="Subtitle"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("150.05–153 MHz射电天文限制、406.1–420 MHz共用结论及典型隔离距离案例"); set_run_font(r,11.5,False,GREY)
doc.add_paragraph()
callout(doc,"文档用途","本文件将前述答复整理为可直接引用和讲解的Word文档，重点给出原文PDF页码、报告内部印刷页码、对应章节与表格，以及相关计算过程和适用范围。")
doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("依据文件：Report ITU-R SA.2427-0（09/2018），共256个PDF页面"); set_run_font(r,10,False,GREY)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("整理日期：2026年9月22日"); set_run_font(r,9.5,False,GREY)
doc.add_page_break()

h = sec.header.paragraphs[0]; h.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = h.add_run("ITU-R SA.2427-0 原文页码定位与分析说明"); set_run_font(r,8.5,False,GREY)
add_page_number(sec.footer.paragraphs[0])

heading(doc,"1. 快速定位表",1)
body(doc,"下表给出最需要查看或截图的原文位置。由于PDF包含封面和前言，PDF阅读器页码通常比报告正文印刷页码大2页。")
table(doc,["内容","PDF阅读器页码","报告印刷页码","章节/表格"],[
("150.05–153 MHz最终汇总结论","第249页","第247页","第10章，表10-2"),
("射电天文详细分析全文","第125–131页","第123–129页","第9.11节，Study 9.11"),
("射电天文保护门限","第125页","第123页","表9.11-1"),
("地球站最小耦合损耗计算","第128–129页","第126–127页","表9.11-4"),
("697 km与560 km隔离距离","第129页","第127页","第9.11.2.1节"),
("Study 9.11总结","第130–131页","第128–129页","第9.11.5节"),
("406.1–420 MHz最终汇总结论","第251页","第249页","第10章，表10-2"),
("陆地移动业务详细分析","第69–76页","第67–74页","第9.4节"),
("固定业务详细分析","第76–81页","第74–79页","第9.5节"),
("国际空间站/空间研究分析","第113–124页","第111–122页","第9.10节"),
],[6.5,3.0,3.0,4.0])

heading(doc,"2. 页码换算说明",1)
body(doc,"这份PDF有封面和前言，因此存在两套页码：一套是PDF阅读器从封面开始计数的物理页码，另一套是页面顶部或底部印刷的报告正文页码。")
callout(doc,"换算示例","PDF第129页对应报告正文印刷第127页。查找原文时，建议优先使用PDF阅读器页码；在技术报告中引用时，可同时标注两套页码。",LIGHT_BLUE)

heading(doc,"3. 150.05–153 MHz射电天文限制结论",1)
heading(doc,"3.1 最终汇总结论的位置",2)
body(doc,"最终结论位于PDF第249页、报告印刷第247页的第10章表10-2。表中150.05–153 MHz一行对地对空（E-s）和空对地（s-E）均标为“No”，并将依据指向Study 9.11（RAS）。")
table(doc,["频段","地对空 E-s","空对地 s-E","原文依据"],[("150.05–153 MHz","No","No","Study 9.11，Radio Astronomy Service")],[4.5,3.2,3.2,7.0],MID_BLUE)

heading(doc,"3.2 详细分析的位置与范围",2)
body(doc,"完整分析集中在PDF第125–131页（报告第123–129页）的第9.11节。该节依次处理射电天文保护门限、卫星空对地下行、地球站地对空上行、单源与聚合效应以及带外/杂散发射。")
table(doc,["分析内容","PDF页码","报告页码"],[
("RAS保护门限，表9.11-1","125","123"),
("单颗NGSO卫星空对地下行分析","125–128","123–126"),
("空对地下行聚合干扰讨论","128","126"),
("地球站地对空单源最小耦合损耗，表9.11-4","128–129","126–127"),
("697 km和560 km隔离距离","129","127"),
("空对地下行带外/杂散发射","129","127"),
("地对空上行带外/杂散发射","130","128"),
("Study 9.11总结","130–131","128–129"),
],[11.0,3.2,3.2])

heading(doc,"3.3 射电天文保护门限",2)
body(doc,"表9.11-1采用ITU-R RA.769-2给出的射电天文保护门限。由于射电天文接收的是极弱宇宙射电信号，其允许干扰功率远低于一般通信或雷达系统。")
table(doc,["RAS频段","输入干扰功率门限","pfd门限","频谱pfd门限"],[
("150.05–153 MHz","−199 dBW","−194 dB(W/m²)","−259 dB(W/(m²·Hz))"),
("406.1–410 MHz","−203 dBW","−189 dB(W/m²)","−255 dB(W/(m²·Hz))"),
],[4.3,4.2,4.2,5.0])

heading(doc,"3.4 空对地：卫星发射干扰射电天文站",2)
body(doc,"报告分别对300 km和1 000 km轨道高度、5°低仰角和90°天顶方向、152 MHz和408 MHz场景进行单干扰源计算。传播损耗考虑自由空间损耗、大气气体衰减和电离层损耗。")
body(doc,"表9.11-2和表9.11-3表明，单颗卫星空对地发射已经超过射电天文保护门限约47–71 dB。多卫星epfd聚合研究尚未完成，但报告预计聚合干扰会进一步增大，因此空对地同频共用被判定为不可行。")

heading(doc,"4. 射电天文隔离的典型量化案例",1)
callout(doc,"核心定位","案例从PDF第128页（报告第126页）的表9.11-4开始；最具代表性的697 km和560 km数值位于PDF第129页（报告第127页）。")
heading(doc,"4.1 计算对象和假设",2)
for x in [
"干扰源为单个短期任务NGSO卫星地球站的地对空发射。",
"地球站发射功率17 dBW，信号带宽25 kHz。",
"射电天文站位于地球站天线旁瓣方向，朝向RAS站的地球站天线增益假设为6 dBi。",
"地球站活动因子100%，天线高度1.5 m。",
"采用ITU-R P.452-16传播模型，包含自由空间、光滑地球绕射、对流层散射和地物杂波衰减。",
"大气单位距离衰减假设为0 dB/km。",
]: bullet(doc,x)
heading(doc,"4.2 量化结果",2)
table(doc,["地球站场景","对应RAS频段宽度","所需隔离距离"],[
("152 MHz地球站 → 150.05–153 MHz RAS","2.95 MHz","697 km"),
("408 MHz地球站 → 406.1–410 MHz RAS","3.90 MHz","560 km"),
],[9.0,4.0,4.0],ORANGE)
body(doc,"Study 9.11总结中将697 km概括为“约700 km”。报告认为，如此大的地理隔离距离使同频共用很难实现。")
body(doc,"该节没有继续完成多地球站聚合计算，因为单个地球站已经产生很高的超限量；进一步聚合只会使结果更严重。")
heading(doc,"4.3 邻频和带外结论",2)
body(doc,"报告利用SA.2426提供的实测杂散发射数据分析带外兼容性，并提出空对地方向的新划分至少应距离两个RAS频段边缘约1.5 MHz。该数值仅是单干扰源估计，真实保护带宽仍需epfd动态聚合仿真确定。")

heading(doc,"5. 406.1–420 MHz“地对空和空对地均不可行”结论",1)
heading(doc,"5.1 最终汇总表的位置",2)
body(doc,"结论位于PDF第251页、报告印刷第249页的表10-2。406.1–420 MHz一行的带宽为13.9 MHz，地对空和空对地均标为“No”。")
table(doc,["频段","带宽","地对空 E-s","空对地 s-E"],[("406.1–420 MHz","13.9 MHz","No","No")],[5.0,4.0,4.0,4.0],MID_BLUE)
heading(doc,"5.2 该结论由多项研究共同支撑",2)
body(doc,"原文不是用一个统一模型直接得出整个406.1–420 MHz的结论，而是分别分析现有业务，再在表10-2中综合归纳。")
table(doc,["受保护业务/系统","覆盖频段或频点","详细研究","PDF页码","主要结论"],[
("COSPAS-SARSAT生命安全系统","406–406.1 MHz及邻频","第7节","约30–32","406.1 MHz附近需保护带并避免新空间操作业务发射"),
("陆地移动业务","406.1–420 MHz","Study 9.4","69–76","卫星下行和地球站上行均不可行"),
("固定业务","406.1–420 MHz","Study 9.5","76–81","卫星下行和地球站上行均不可行"),
("空间研究/ISS系统","414.2、417.1 MHz","Study 9.10","113–124","研究的双向场景超过保护准则"),
("射电天文业务","406.1–410 MHz","Study 9.11","125–131","同频共用困难/不可行"),
],[4.5,3.6,3.0,2.4,5.0])
heading(doc,"5.3 陆地移动业务：Study 9.4",2)
body(doc,"Study 9.4位于PDF第69–76页。报告采用ITU-R M.1808中的移动基站和移动台参数，以I/N = −6 dB为一般保护准则；对公共保护和灾害救援等高保护需求应用，可采用I/N = −10 dB。")
body(doc,"单颗NGSO卫星下行对基站和移动台的保护门限超限最高超过39 dB；考虑多颗同频卫星后，聚合干扰更严重。对于地球站地对空发射，所需360°排斥区过大，因而两个方向均不可行。")
heading(doc,"5.4 固定业务：Study 9.5",2)
body(doc,"Study 9.5位于PDF第76–81页。空对地下行中，单颗NGSO卫星对固定业务保护门限的超限约为36 dB，多星聚合后会更严重。地对空方向则需要很大的地理隔离距离，因而报告判定两个方向均不可行。")
heading(doc,"5.5 空间研究/国际空间站：Study 9.10",2)
body(doc,"Study 9.10位于PDF第113–124页，主要研究国际空间站空间对空间通信系统的414.2 MHz主频和417.1 MHz备用频率。")
for x in [
"建立8类双向干扰场景，包括ISS电台干扰NGSO卫星/地球站，以及NGSO卫星/地球站干扰ISS接收机。",
"采用300颗卫星和300个一一对应地球站。",
"动态仿真持续7天，时间步长20 s。",
"轨道传播考虑J2摄动和WGS-84地球扁率。",
"统计最小安全距离、排斥半径、干扰事件时间百分比和接收功率。",
]: bullet(doc,x)
body(doc,"在所研究的具体系统参数和运行假设下，多个场景超过保护准则，报告认为在410–420 MHz内这些场景不能共用。")
heading(doc,"5.6 射电天文约束的实际覆盖范围",2)
body(doc,"射电天文分析覆盖的是406.1–410 MHz，而不是整个406.1–420 MHz。整个频段的“No/No”结论，是由射电天文、陆地移动、固定、COSPAS-SARSAT和ISS空间研究等多项独立研究共同支撑的综合结论。")

heading(doc,"6. 结论的适用范围与正确表述",1)
callout(doc,"适用范围提醒","表10-2中的“No”针对SA.2426定义的短期任务非GSO卫星及地球站典型参数，并基于SA.2427采用的传播、天线、运行与保护准则假设。它不应被扩大解释为任何参数、任何空间操作系统在相关频段都绝对不可能工作。","FFF2CC",ORANGE)
body(doc,"更严谨的表述为：按SA.2426给出的典型参数，在SA.2427所研究的条件下，150.05–153 MHz主要受射电天文极低保护门限限制；406.1–420 MHz则由移动、固定、射电天文、COSPAS-SARSAT和ISS空间研究等多项独立兼容性分析共同支持No/No结论。")

heading(doc,"7. 推荐截图和引用位置",1)
for x in [
"PDF第125页：表9.11-1，射电天文保护门限。",
"PDF第128页：表9.11-4开始，地球站最小耦合损耗计算。",
"PDF第129页：697 km和560 km隔离距离，是最适合展示的量化案例。",
"PDF第130–131页：Study 9.11总结，含47–71 dB超限、约700/560 km隔离和1.5 MHz保护带结论。",
"PDF第249页：表10-2中150.05–153 MHz的No/No汇总。",
"PDF第251页：表10-2中406.1–420 MHz的No/No汇总及对应研究索引。",
]: numbered(doc,x)

heading(doc,"8. 一句话总结",1)
callout(doc,"总结","150.05–153 MHz的结论有完整的保护门限、单星下行、P.452最小耦合损耗和地理隔离分析；最典型的697 km与560 km案例位于PDF第129页。406.1–420 MHz的No/No结论位于PDF第251页，是移动、固定、射电天文、COSPAS-SARSAT和ISS空间研究等多项分析的综合汇总。")

core = doc.core_properties
core.title = "ITU-R SA.2427-0原文页码定位与分析说明"
core.subject = "150.05–153 MHz射电天文限制、406.1–420 MHz共用结论及典型隔离距离案例"
core.author = "OpenAI"
core.keywords = "ITU-R, SA.2427-0, 射电天文, 共用分析, 隔离距离"
core.comments = "依据用户提供的Report ITU-R SA.2427-0整理。"

doc.save(OUT)
print(OUT)
